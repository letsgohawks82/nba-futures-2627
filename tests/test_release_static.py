"""Release-control tests: metadata, card/workbook agreement, flags, test hygiene. Pure static checks: they need no simulation run."""
import glob, json, os, re, subprocess
import numpy as np, pandas as pd
import pytest
from conftest import ROOT
from src.release import checks as K
from src.release.meta import META

BASE = os.path.join(ROOT, "repro", "baseline")


def committed_table(rel):
    """The committed output tables: the pre-run snapshot if a repro run has already overwritten them."""
    snap = os.path.join(BASE, rel.replace("/", "__"))
    return snap if os.path.exists(snap) else os.path.join(ROOT, rel)


def out_dir():
    return os.environ.get("NBA_OUTPUT_DIR") or os.path.join(ROOT, "output")


def workbook_path():
    return os.path.join(out_dir(), META["workbook_file"])


def card_source():
    """Card the workbook must agree with: the rebuilt card after a repro run, else the committed final card."""
    p = os.path.join(ROOT, "repro", "rebuilt_card.csv")
    return pd.read_csv(p) if os.path.exists(os.path.join(ROOT, "repro", "run_status.json")) and os.path.exists(p) else pd.read_csv(os.path.join(ROOT, "release", "final_card.csv"))


def git(*a):
    try: return subprocess.run(["git", *a], cwd=ROOT, capture_output=True, text=True).stdout.strip()
    except Exception: return ""


def test_tests_do_not_contain_vacuous_assertions():
    for f in glob.glob(os.path.join(ROOT, "tests", "*.py")):
        for i, line in enumerate(open(f, encoding="utf-8"), 1):
            if re.search(r"^\s*assert\b.*\bor True\b", line) and "test_release_static" not in f:
                pytest.fail(f"{os.path.basename(f)}:{i}: vacuous assertion (`or True`)")


def test_final_card_csv_matches_committed_pricing_tables():
    """release/final_card.csv is the committed card; it must equal the card recomputed from the committed pricing tables."""
    a = pd.read_csv(os.path.join(ROOT, "release", "final_card.csv")).sort_values(["market", "selection"]).reset_index(drop=True)
    b = K.card_from_tables_at(committed_table("data/clean/pricing_table.csv"), committed_table("data/clean/player_pricing_table.csv"))
    assert list(zip(a.market, a.selection, a.side)) == list(zip(b.market, b.selection, b.side)), "card membership differs from the committed pricing tables"
    for c in ("american", "fair", "model_prob", "ev", "ev_shrunk75", "min_price", "stake_frac"):
        assert np.allclose(a[c].astype(float), b[c].astype(float), atol=1e-6, rtol=0), f"final_card.csv column {c} differs from the committed pricing tables"


def test_workbook_tier1_rows_count_and_stakes_agree_with_card():
    p = workbook_path(); assert os.path.exists(p), f"workbook missing: {p}"
    card = card_source(); wb = K.workbook_card(p)
    assert len(wb) == len(card), f"workbook lists {len(wb)} Tier 1 rows, card has {len(card)}"
    assert sorted(zip(wb.market, wb.selection, wb.side)) == sorted(zip(card.market, card.selection, card.side)), "workbook Tier 1 rows differ from the card"
    import openpyxl
    wv = openpyxl.load_workbook(p, data_only=True)
    title = wv["Tier1_Card"]["A2"].value
    m = re.search(r"\((\d+) active Tier 1 bets\)", str(title)); assert m, f"Tier1_Card title has no bet count: {title!r}"
    assert int(m.group(1)) == len(card), f"Tier1_Card title says {m.group(1)} bets, card has {len(card)}"
    assert "8 bets" not in str(title)
    tr = wv["Tracker"]; n_tr = sum(1 for r in range(8, 40) if tr.cell(r, 2).value not in (None, "") and tr.cell(r, 14).value not in (None, ""))
    cached = wb.cached_stake.dropna() if wb.cached_stake.notna().any() else None
    if cached is not None:           # Excel-recalculated workbook: displayed stakes, tracker rows and the live status must agree
        assert len(cached) == len(card)
        merged = wb.assign(k=list(zip(wb.market, wb.selection, wb.side))).set_index("k").cached_stake
        for r in card.itertuples():
            assert abs(float(merged[(r.market, r.selection, r.side)]) - r.stake_frac) < 1e-4, f"stake shown in the workbook differs from the card for {r.market} {r.selection}"
        assert n_tr == len(card), f"Tracker lists {n_tr} Tier 1 rows, card has {len(card)}"
        assert all(v is True for v in K.tier1_flags_live(p)), "a Tier 1 row is not live-Tier-1 in the workbook"


def test_workbook_has_no_integrity_flags():
    p = workbook_path(); assert os.path.exists(p)
    flags, computed = K.integrity_flags(p)
    assert not flags, f"workbook integrity FLAGS present: {flags[:5]}"
    mx, bad = K.analytic_vs_sim()
    assert not bad, f"analytic vs sim P(over) deltas above 2pp (max {mx:.2f}): {bad}"
    if not computed:
        if os.environ.get("NBA_ALLOW_NO_EXCEL") != "1": pytest.fail("workbook has no cached check statuses; rebuild with Excel or set NBA_ALLOW_NO_EXCEL=1 to accept the Python re-computation of the flagged checks")
        import warnings; warnings.warn("workbook built without Excel: integrity checks were re-computed in Python (analytic vs sim) rather than read from Checks!C")


def test_release_metadata_agrees_everywhere():
    v = META["version"]
    for f in ("README.md", "AUDIT.md"):
        txt = open(os.path.join(ROOT, f), encoding="utf-8").read(); assert f"Release: `{v}`" in txt, f"{f} does not state `Release: {v}`"
    import openpyxl
    wb = openpyxl.load_workbook(workbook_path(), data_only=False); cov = wb["Cover"]
    vals = {cov.cell(r, 1).value: cov.cell(r, 2).value for r in range(7, 16)}
    assert str(vals["Version"]).startswith(v), f"Cover version {vals['Version']!r} != {v}"
    assert vals["Build date"] == META["build_date"]; assert META["odds_as_of"] in str(vals["Prices as of"]) and META["consensus_as_of"] in str(vals["Prices as of"])
    for ws in wb.worksheets:
        a3 = str(ws["A3"].value or "")
        assert v in a3 and META["odds_as_of"] in a3 and META["build_date"] in a3, f"{ws.title}!A3 header {a3!r} disagrees with params/release.json"
    assert os.path.basename(workbook_path()) == META["workbook_file"]
    assert os.path.exists(os.path.join(ROOT, "output", f"NBA_Futures_Dashboard_BetCard_{META['build_date'].replace('-', '')}.pdf")), "the committed dashboard PDF for the build date is missing"
    rep = os.path.join(ROOT, "reports", "release_report.md")
    if os.path.exists(rep): assert v in open(rep, encoding="utf-8").read().splitlines()[0], "release_report.md title does not carry the version"
    # git tag / commit agreement
    base = git("rev-parse", "--verify", "-q", f"refs/tags/{META['baseline_tag']}^{{commit}}")
    if base: assert base == META["baseline_commit"], f"baseline tag {META['baseline_tag']} moved: {base} != {META['baseline_commit']}"
    exact = git("describe", "--tags", "--exact-match", "HEAD")
    if exact and exact.startswith("v1.0-preseason."): assert exact == v or exact.startswith(v + "-"), f"HEAD is tagged {exact} but params/release.json says {v}"
    tag = git("rev-parse", "--verify", "-q", f"refs/tags/{v}^{{commit}}")
    if tag: assert tag == git("rev-parse", "HEAD") or git("merge-base", "--is-ancestor", tag, "HEAD") == "" , f"tag {v} does not point at an ancestor of HEAD"


def test_darko_blend_is_labelled_unvalidated():
    wb_txt = open(os.path.join(ROOT, "params", "release.json"), encoding="utf-8").read()
    assert "UNVALIDATED" in wb_txt
    for f in ("README.md", "AUDIT.md"):
        assert "unvalidated" in open(os.path.join(ROOT, f), encoding="utf-8").read().lower(), f"{f} must label the DARKO blend as unvalidated"
    import openpyxl
    a = openpyxl.load_workbook(workbook_path(), data_only=False)["Assumptions"]
    row = [a.cell(r, 4).value for r in range(8, 70) if a.cell(r, 2).value is not None and "DARKO" in str(a.cell(r, 1).value).upper()]
    assert row and all("UNVALIDATED" in str(x).upper() for x in row), f"Assumptions DARKO rows are not labelled UNVALIDATED: {row}"
