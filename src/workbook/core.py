"""Workbook core tabs: Assumptions, Raw_*, Overrides, Odds_Input, M_DeVig."""
import json, os
import numpy as np, pandas as pd
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.utils import get_column_letter as CL
from .styles import *

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
C = lambda f: os.path.join(ROOT, "data", "clean", f)
REG = {}          # registry of cross-sheet references
N_ITER = 32


def name(wb, nm, ref):
    wb.defined_names[nm] = DefinedName(nm, attr_text=ref)


# ----------------------------------------------------------------------------- Assumptions
def build_assumptions(wb):
    ws = wb.create_sheet("Assumptions"); sheet_header(ws, "Assumptions - every knob in one table")
    prov = json.load(open(os.path.join(ROOT, "params/provenance.json"))); gm = json.load(open(C("game_model_fit.json")))
    fit = json.load(open(C("ratings_fit_2027.json"))); L = fit["team_layer"]; meta = json.load(open(C("sim_run_meta.json")))
    pf = json.load(open(C("playoff_adj_fit.json")))
    put(ws, "A5", "Scenario selector (Bull / Base / Bear)", "text", bold=True); put(ws, "B5", "Base", "input")
    dv = DataValidation(type="list", formula1='"Bull,Base,Bear"', allow_blank=False); ws.add_data_validation(dv); dv.add("B5"); name(wb, "A_Scenario", "Assumptions!$B$5")
    put(ws, "C5", "Bull = lower sigma_team / higher DARKO weight / no tank; Bear = higher sigma_team / lower DARKO weight / upper-CI tank. Reprices the analytic layer live; sims must be rerun.", "text", italic=True)
    headers(ws, 7, ["Name", "Value (active)", "Units", "Description", "Provenance", "Default", "Bull", "Base", "Bear"])
    sc = {}  # name -> (bull, base, bear)
    rows = [  # (name, key, value, units, desc, prov, scenario(bull,base,bear)|None, fmt, kind)
        ("HCA", "A_hca", gm["hca"], "pts", "Home-court advantage (regular season), margin points", "params/provenance.json:game_hca (2021-22..2025-26)", None, "num2", "input"),
        ("sigma_game", "A_sigma_game", gm["sigma_game"], "pts", "Game margin standard deviation", "game_sigma_game", None, "num2", "input"),
        ("sigma_team", "A_sigma_team", 3.5, "SRS pts", "True-strength shock SD (season) - LOSO CV with mean-matched sims, availability draws on", "sigma_team_srs (G5 refit)", (3.15, 3.5, 3.85), "num2", "input"),
        ("prior_blend_weight (coef on prior SRS)", "A_b2", L["beta"][2], "coef", "Team layer: coefficient on prior-season SRS", "team_layer_beta[2]", None, "num3", "input"),
        ("team_layer_const", "A_b0", L["beta"][0], "coef", "Team layer: intercept", "team_layer_beta[0]", None, "num3", "input"),
        ("team_layer_rating_coef", "A_b1", L["beta"][1], "coef", "Team layer: coefficient on minutes-weighted player rating sum", "team_layer_beta[1]", None, "num3", "input"),
        ("team_layer_continuity_coef", "A_b3", L["beta"][3], "coef", "Team layer: coefficient on continuity x prior SRS", "team_layer_beta[3]", None, "num3", "input"),
        ("k_wins", "A_kwins", L["k_wins"], "wins/SRS pt", "Wins per SRS point (wins = 41 + k x SRS)", "k_wins_per_srs", None, "num3", "input"),
        ("DARKO_weight (player-metric blend)", "A_wdarko", 0.5, "share", "Production DARKO vs box-prior blend (documented rule, untested)", "w_darko (G2 addendum); engine input", (0.7, 0.5, 0.3), "num2", "input"),
        ("box_prior_weight", "A_wbox", 0.5, "share", "Equals 1 minus DARKO weight (engine input)", "derived", None, "num2", "formula"),
        ("aging_curve_on", "A_aging", 1, "0/1", "Aging curve in player projections (engine input; informational)", "ratings fit", None, "int", "input"),
        ("rookie_prior_weight", "A_rookie", 1.0, "share", "Rookie draft-slot + college translation model weight (engine input; informational)", "rookie_college_model_beta", None, "num2", "input"),
        ("tank_on", "A_tank_on", 1, "0/1", "3-2-1 tank module on (engine input)", "G3", (0, 1, 1), "int", "input"),
        ("tank_kappa", "A_kappa", meta["par"]["tank_kappa"], "net-rating pts/unit", "Post-deadline effort suppression per unit normalized incentive (CI includes 0)", "tank_kappa (G3)", (0.0, meta["par"]["tank_kappa"], 1.57), "num4" if False else "0.0000", "input"),
        ("tank_deterrent", "A_deterrent", 0.0, "share", "Commissioner deterrent (1 = no tanking)", "tank_deterrent (assumption)", None, "num2", "input"),
        ("playoff_adj_lambda", "A_lambda", pf["lambda_po"], "share", "Playoff minutes-concentration rating adjustment multiplier (fit 0.58 [0.32,0.80])", "po_lambda", None, "num3", "input"),
        ("playoff_injury_prob", "A_poinj", 0.07, "prob", "Per top-8 player chance of missing the postseason (assumption)", "po_inj_prob", None, "num2", "input"),
        ("lottery_n_draws", "A_ndraws", 16, "draws", "Lottery draws (all 16 picks drawn; NBA.com)", "assumptions.json", None, "int", "input"),
        ("sim_N", "A_nsims", meta["n_sims"], "seasons", "Simulated seasons in the last run", "sims/run_meta.json", None, "int", "input"),
        ("cup_semis_neutral", "A_cupneutral", 0, "0/1", "NBA.com: semis at higher seed home arena", "assumptions.json", None, "int", "input"),
        ("cup_nonqualifier_cross_conf_games", "A_cupcross", 2, "games", "NBA.com: 2 of the 22 non-qualifier games are cross-conference", "assumptions.json", None, "int", "input"),
        ("tpg_min_games_pct", "A_tpg", 0.70, "share", "Turnover-leader qualifier (58-game rule; UNCONFIRMED with the book)", "assumptions.json", None, "num2", "input"),
        ("games_65_rule_on", "A_65", 1, "0/1", "65-game eligibility rule for awards (G4)", "SPEC", None, "int", "input"),
        ("devig_method", "A_devig", "Shin", "list", "De-vig method for fair probabilities (Multiplicative / Power / Shin)", "SPEC", None, "text", "input"),
        ("bankroll", "A_bankroll", None, "$", "Bankroll (blank input; stakes show '-' until entered)", "user input", None, "money", "input"),
        ("kelly_fraction", "A_kelly", 0.25, "share", "Fractional Kelly multiplier", "SPEC", None, "num2", "input"),
        ("min_EV", "A_minev", 0.03, "per $1", "Minimum EV to qualify for the Bet Card", "SPEC", None, "pct", "input"),
        ("CI_level", "A_cilevel", 0.80, "prob", "Confidence level of the edge CI (engine output is computed at 80%)", "SPEC", None, "pct", "input"),
        ("max_bet", "A_maxbet", 1000, "$", "Max single bet", "SPEC / book MAX", None, "money", "input"),
        ("max_total_exposure", "A_maxexp", 0.25, "share of bankroll", "Max total exposure", "SPEC", None, "pct", "input"),
        ("staking_shrink_to_market", "A_shrink", 0.50, "share", "Model-to-market shrink (0 = pure model). 0.5 after the G5 miscalibration diagnosis (model-vs-market wins RMSE > 4); also applied to the edge test; the model itself is never shrunk", "SPEC default 0.25; raised to 0.5 (G5 report)", None, "num2", "input"),
        ("zero_sum_filter_on", "A_zs_on", 1, "0/1", "WINS bets must survive rescaling model win means to the book's sum of lines (removes edges that are only market inflation)", "G5 report", None, "int", "input"),
        ("longshot_threshold", "A_longthr", 0.05, "prob", "Selections with fair probability below this must survive shrink + sigma_team -10% (longshot control)", "G5 report", None, "pct", "input"),
        ("low_continuity_threshold", "A_contthr", 0.55, "share", "Teams with roster continuity below this need wins/playoffs/division edges to survive the stronger shrink", "G5 report", None, "num2", "input"),
        ("low_continuity_shrink", "A_contshrink", 0.75, "share", "Shrink applied to the edge test for low-continuity teams", "G5 report", None, "num2", "input"),
        ("zero_sum_shift", "A_zs_delta", "=(SUMIFS(Odds_Input!$E$8:$E$694,Odds_Input!$B$8:$B$694,\"WINS\",Odds_Input!$D$8:$D$694,\"OVER\")-A_leaguew)/30", "wins/team", "Book sum of win lines minus 1,230, per team (live)", "derived from Odds_Input", None, "num3", "formula"),
        ("playoffs_definition", "A_podef", "Bracket", "list", "PLAYOFFS market definition: Bracket (16-team playoffs; matches book prices) or Top10 (play-in counts; spec reading)", "UNCONFIRMED - see G5 report", None, "text", "input"),
        ("playoffs_rule_confirmed", "A_poconf", 1, "0/1", "PLAYOFFS = 16-team playoff field (top-8, including via play-in): user-confirmed from implied sums (~8.7/conf). CONFIRM WITH BOOK. Set 0 to block PLAYOFFS bets", "G5 report", None, "int", "input"),
        ("analytic_sd_scale", "A_sdscale", 0.94, "ratio", "Analytic layer: SD scale vs delta-method SD (fit to sim SDs)", "G5: fit of analytic layer to the engine", None, "num2", "input"),
        ("games_per_season", "A_games", 82, "games", "Games per team (structural constant)", "SPEC", None, "int", "input"),
        ("team_minutes_per_season", "A_teamminutes", 19680, "minutes", "240 minutes x 82 games (structural constant)", "SPEC", None, "int", "input"),
        ("pace_blend_prior", "A_paceblend", 0.7, "share", "Pace projection = share x prior pace + (1 - share) x league mean (assumption, untested)", "provenance pace_blend", None, "num2", "input"),
        ("playin_slots", "A_playin_slots", 4, "slots", "Play-in participants per conference (finish 7-10); target sum for PLAYIN fair probabilities", "NBA rules", None, "int", "input"),
        ("lg_ortg", "A_lgortg", fit["lg_ortg"], "pts/100", "League average ORtg 2025-26 (anchor for ORtg/DRtg)", "ratings_fit_2027.json", None, "num2", "input"),
        ("league_wins_total", "A_leaguew", 1230, "wins", "Zero-sum constraint on projected wins", "SPEC", None, "int", "input"),
    ]
    r = 8
    for nm, key, val, units, desc, pv, scn, fmt, kind in rows:
        put(ws, f"A{r}", nm, "text")
        if scn:
            put(ws, f"G{r}", scn[0], "input", fmt); put(ws, f"H{r}", scn[1], "input", fmt); put(ws, f"I{r}", scn[2], "input", fmt)
            put(ws, f"B{r}", f'=CHOOSE(MATCH(A_Scenario,{{"Bull","Base","Bear"}},0),G{r},H{r},I{r})', "formula", fmt)
        else:
            if key == "A_wbox":
                put(ws, f"B{r}", f"=1-A_wdarko", "formula", fmt)
            elif isinstance(val, str) and val.startswith("="):
                put(ws, f"B{r}", val, "formula", fmt)
            elif val is None:
                put(ws, f"B{r}", None, "input", fmt)
            else:
                put(ws, f"B{r}", val, "input", fmt)
            for col in "GHI":
                put(ws, f"{col}{r}", "-", "text", align="right")
        put(ws, f"C{r}", units, "text"); put(ws, f"D{r}", desc, "text"); put(ws, f"E{r}", pv, "text")
        put(ws, f"F{r}", scn[1] if scn else ("-" if (val is None or (isinstance(val, str) and val.startswith("="))) else val), "text" if (val is None or (isinstance(val, str) and val.startswith("="))) else "input", fmt if (val is not None and not (isinstance(val, str) and val.startswith("="))) else None)
        name(wb, key, f"Assumptions!$B${r}"); REG[key] = f"Assumptions!$B${r}"
        r += 1
    dv2 = DataValidation(type="list", formula1='"Multiplicative,Power,Shin"'); ws.add_data_validation(dv2); dv2.add(f"B{[i for i,x in enumerate(rows) if x[1]=='A_devig'][0]+8}")
    dv3 = DataValidation(type="list", formula1='"Bracket,Top10"'); ws.add_data_validation(dv3); dv3.add(f"B{[i for i,x in enumerate(rows) if x[1]=='A_podef'][0]+8}")
    # sim hash key: concatenation of sim-relevant knobs (same string is stored by Python when the sims run)
    r += 1
    keys = ["A_sigma_team", "A_hca", "A_sigma_game", "A_wdarko", "A_kappa", "A_deterrent", "A_tank_on", "A_lambda", "A_poinj", "A_ndraws", "A_cupneutral", "A_cupcross", "A_b0", "A_b1", "A_b2", "A_b3", "A_kwins"]
    put(ws, f"A{r}", "Sim assumptions key (current)", "text", bold=True)
    expr = '&"|"&'.join(('TEXT(IF(A_cupcross>0,1,0),"0.0000")' if k == "A_cupcross" else f'TEXT({k},"0.0000")') for k in keys)
    put(ws, f"B{r}", "=" + expr, "formula"); name(wb, "A_HashCurrent", f"Assumptions!$B${r}"); REG["hash_current"] = f"Assumptions!$B${r}"
    REG["hash_keys"] = keys
    put(ws, f"D{r}", "Compared with the key stored with the last sim run (M_Sims); a mismatch flags sim-dependent cells STALE - rerun sims", "text", italic=True)
    widths(ws, {"A": 38, "B": 16, "C": 18, "D": 70, "E": 40, "F": 14, "G": 12, "H": 12, "I": 12}); ws.freeze_panes = "A8"
    return ws


def hash_key_string(values):
    return "|".join(f"{v:.4f}" for v in values)


# ----------------------------------------------------------------------------- Raw tabs
def build_raw(wb):
    T = pd.read_csv(C("team_projection_2027.csv")).sort_values("abbr").reset_index(drop=True)
    P = pd.read_csv(C("player_projection_2027.csv"))
    S = pd.read_csv(C("schedule_games.csv"))
    ws = wb.create_sheet("Raw_Teams"); sheet_header(ws, "Raw_Teams - team inputs (clean data)")
    from ..data.teams import TEAMS
    cols = ["Team", "Conf", "Division", "Prev SRS", "Prev wins", "Continuity", "Prev pace", "Prev ORtg", "Prev DRtg", "TeamSum O (BPM)", "TeamSum D (BPM)", "Python pred SRS", "Python pred wins", "Python teamsum"]
    headers(ws, 7, cols)
    for i, r in T.iterrows():
        row = 8 + i
        put(ws, f"A{row}", r.abbr, "input"); put(ws, f"B{row}", TEAMS[r.abbr][1], "input"); put(ws, f"C{row}", TEAMS[r.abbr][2], "input")
        for col, v, f in [("D", r.prev_srs, "num2"), ("E", r.prev_wins, "int"), ("F", r.continuity, "num3"), ("G", r.prev_pace, "num1"), ("H", r.prev_ortg, "num1"), ("I", r.prev_drtg, "num1"),
                          ("J", r.ts_o, "num3"), ("K", r.ts_d, "num3"), ("L", r.pred_srs, "num3"), ("M", r.pred_wins, "num2"), ("N", r.teamsum, "num3")]:
            put(ws, f"{col}{row}", float(v), "input", f)
    REG["raw_teams"] = dict(first=8, last=7 + len(T), n=len(T))
    widths(ws, {"A": 10, "B": 8, "C": 12}); finish(ws)
    ws = wb.create_sheet("Raw_Players"); sheet_header(ws, "Raw_Players - 2026-27 roster projections (clean data)")
    headers(ws, 7, ["Team", "Player", "bbref id", "Age", "Rating (blended)", "Projected minutes", "Projected mpg", "Projected games", "No NBA history", "DARKO DPM"])
    P = P.sort_values(["abbr", "proj_min"], ascending=[True, False]).reset_index(drop=True)
    for i, r in P.iterrows():
        row = 8 + i
        put(ws, f"A{row}", r.abbr, "input"); put(ws, f"B{row}", r["name"], "input"); put(ws, f"C{row}", str(r.bbref_id), "input"); put(ws, f"D{row}", float(r.age), "input", "num1")
        put(ws, f"E{row}", float(r.rating), "input", "num3"); put(ws, f"F{row}", float(r.proj_min), "input", "num1"); put(ws, f"G{row}", float(r.proj_mpg), "input", "num1"); put(ws, f"H{row}", float(r.proj_gp), "input", "num1")
        put(ws, f"I{row}", int(bool(r.no_history)), "input", "int"); put(ws, f"J{row}", float(r.dpm) if pd.notna(r.dpm) else "-", "input", "num1")
    REG["raw_players"] = dict(first=8, last=7 + len(P), n=len(P))
    widths(ws, {"A": 8, "B": 28, "C": 14}); finish(ws, label_cols=("B",))
    ws = wb.create_sheet("Raw_Schedule"); sheet_header(ws, "Raw_Schedule - 2026-27 scheduled games (80 per team; 30 Cup-dependent games are simulated)")
    headers(ws, 7, ["Game #", "Date", "Start (ET)", "Away", "Home", "Cup group game", "Arena"])
    for i, r in S.iterrows():
        row = 8 + i
        put(ws, f"A{row}", int(r.game_no), "input", "int"); put(ws, f"B{row}", r.date, "input"); put(ws, f"C{row}", r.start_et, "input"); put(ws, f"D{row}", r.away, "input"); put(ws, f"E{row}", r.home, "input")
        put(ws, f"F{row}", int(bool(r.cup_group)), "input", "int"); put(ws, f"G{row}", r.arena, "input")
    widths(ws, {"G": 30}); finish(ws)
    return T, P


# ----------------------------------------------------------------------------- Overrides
def build_overrides(wb, P):
    ws = wb.create_sheet("Overrides"); sheet_header(ws, "Overrides - per-player minutes / games / rating adjustments, per-team rating adjustments (blue = input)")
    put(ws, "A5", "Player overrides (leave blank for no override). Minutes/game and Games replace the projection; rating adjustment is in BPM points per 100 poss.", "text", italic=True)
    headers(ws, 7, ["Player", "Minutes per game override", "Games override", "Rating adjustment", "Note"])
    n = REG["raw_players"]
    dv = DataValidation(type="list", formula1=f"=Raw_Players!$B${n['first']}:$B${n['last']}", allow_blank=True); ws.add_data_validation(dv)
    for r in range(8, 48):
        for col, f in (("A", None), ("B", "num1"), ("C", "int"), ("D", "num2"), ("E", None)):
            put(ws, f"{col}{r}", None, "input", f)
        dv.add(f"A{r}")
    put(ws, "A50", "Team rating adjustments (SRS pts, applied after the team layer, re-centered to zero-sum)", "text", bold=True)
    headers(ws, 51, ["Team", "SRS adjustment", "Note"])
    T = pd.read_csv(C("team_projection_2027.csv")).sort_values("abbr").reset_index(drop=True)
    for i, a in enumerate(T.abbr):
        put(ws, f"A{52+i}", a, "text"); put(ws, f"B{52+i}", 0, "input", "num2"); put(ws, f"C{52+i}", None, "input")
    REG["ov_player"] = (8, 47); REG["ov_team"] = (52, 81)
    widths(ws, {"A": 28, "B": 20, "C": 16, "D": 18, "E": 40}); finish(ws)


# ----------------------------------------------------------------------------- Odds_Input + M_DeVig
def build_odds(wb):
    o = pd.read_csv(os.path.join(ROOT, "data/raw/odds/odds_2026-10-15.csv"))
    ws = wb.create_sheet("Odds_Input"); sheet_header(ws, "Odds_Input - paste area (blue) + parsed book odds; orange fill = imputed/backfilled")
    put(ws, "A5", "Paste/parse book odds here. Command: run.py parse-odds <file>. Source file: odds_2026-10-15.csv (687 rows). Orange = backfilled/imputed (imputed sides are never bet).", "text", italic=True)
    headers(ws, 7, ["Row id", "Market", "Selection", "Side", "Line", "American odds", "Flag"])
    fill_o = PatternFill("solid", fgColor=ORANGE)
    for i, r in o.iterrows():
        row = 8 + i
        put(ws, f"A{row}", i, "input", "int"); put(ws, f"B{row}", r.market, "input"); put(ws, f"C{row}", r.selection, "input"); put(ws, f"D{row}", r.side, "input")
        put(ws, f"E{row}", float(r.line) if pd.notna(r.line) else "-", "input", "num1"); put(ws, f"F{row}", int(r.american), "input", "amer"); put(ws, f"G{row}", r.flag if pd.notna(r.flag) else "-", "input")
        if pd.notna(r.flag):
            for col in "ABCDEFG":
                ws[f"{col}{row}"].fill = fill_o
    n = len(o); REG["odds"] = dict(first=8, last=7 + n, n=n)
    widths(ws, {"A": 8, "B": 16, "C": 30, "D": 10, "E": 10, "F": 14, "G": 22}); finish(ws, label_cols=("C",))
    return o


def build_devig(wb, o):
    ws = wb.create_sheet("M_DeVig"); sheet_header(ws, "M_DeVig - implied probabilities, overround, fair probabilities (Multiplicative / Power / Shin); two-way markets de-vigged as pairs")
    put(ws, "A5", "PLAYIN markets target a sum of 4 (four play-in slots per conference); Shin is undefined for S<>1 so those use the Power method. Solver blocks (right) are grouped, not hidden.", "text", italic=True)
    hdr = ["Row id", "Market", "Selection", "Side", "Line", "American", "Group key", "Decimal", "Implied prob", "Target sum S", "Outcomes in group", "Sum of implied", "Overround", "Fair - Multiplicative", "Fair - Power", "Fair - Shin", "Fair - selected", "Power k", "Shin z"]
    headers(ws, 7, hdr); n = REG["odds"]; f0, f1 = n["first"], n["last"]; first, last = 8, 7 + n["n"]
    # unique keys (python mirror for the solver block)
    keys = []
    for r in o.itertuples():
        k = r.market if r.market not in ("WINS", "PLAYOFFS") else f"{r.market}|{r.selection}" + (f"|{r.line:.1f}" if r.market == "WINS" else "")
        keys.append(k)
    ukeys = list(dict.fromkeys(keys)); SOL0 = 8; sol_first = 8; sol_last = 7 + len(ukeys)
    # column map of the solver block: starts at column 24 (X)
    c0 = 24
    def sc(i): return CL(c0 + i)
    # solver headers
    put(ws, f"{sc(0)}7", "Solver key", "header"); put(ws, f"{sc(1)}7", "S", "header"); put(ws, f"{sc(2)}7", "Sum implied", "header")
    # power: columns 3..: per iteration 4 cols (mid, g, lo, hi) ; initial lo/hi in cols 3,4
    p_start = 3; put(ws, f"{sc(p_start)}7", "Power lo0", "header"); put(ws, f"{sc(p_start+1)}7", "Power hi0", "header")
    def pcol(t, j): return p_start + 2 + 4 * (t - 1) + j       # j: 0 mid,1 g,2 lo,3 hi
    for t in range(1, N_ITER + 1):
        for j, lab in enumerate(["mid", "g", "lo", "hi"]):
            put(ws, f"{sc(pcol(t, j))}7", f"P{t} {lab}", "header")
    s_start = pcol(N_ITER, 3) + 1; put(ws, f"{sc(s_start)}7", "Shin lo0", "header"); put(ws, f"{sc(s_start+1)}7", "Shin hi0", "header")
    def zcol(t, j): return s_start + 2 + 4 * (t - 1) + j
    for t in range(1, N_ITER + 1):
        for j, lab in enumerate(["mid", "g", "lo", "hi"]):
            put(ws, f"{sc(zcol(t, j))}7", f"Z{t} {lab}", "header")
    out_p = zcol(N_ITER, 3) + 1
    put(ws, f"{sc(out_p)}7", "Power k*", "header"); put(ws, f"{sc(out_p+1)}7", "Shin z*", "header")
    KEYR = f"$G${first}:$G${last}"; IMPR = f"$I${first}:$I${last}"
    for i, k in enumerate(ukeys):
        row = sol_first + i
        put(ws, f"{sc(0)}{row}", k, "text"); put(ws, f"{sc(1)}{row}", f'=IF(LEFT({sc(0)}{row},6)="PLAYIN",A_playin_slots,1)', "formula")
        put(ws, f"{sc(2)}{row}", f"=SUMIF({KEYR},{sc(0)}{row},{IMPR})", "formula", "num4" if False else "0.0000")
        put(ws, f"{sc(p_start)}{row}", 0.2, "input"); put(ws, f"{sc(p_start+1)}{row}", 10, "input")
        for t in range(1, N_ITER + 1):
            lo_prev = f"{sc(p_start if t == 1 else pcol(t-1, 2))}{row}"; hi_prev = f"{sc(p_start+1 if t == 1 else pcol(t-1, 3))}{row}"
            mid = f"{sc(pcol(t,0))}{row}"; g = f"{sc(pcol(t,1))}{row}"
            put(ws, mid, f"=({lo_prev}+{hi_prev})/2", "formula", "0.000000")
            put(ws, g, f"=SUMPRODUCT(({KEYR}={sc(0)}{row})*({IMPR}^{mid}))-{sc(1)}{row}", "formula", "0.000000")
            put(ws, f"{sc(pcol(t,2))}{row}", f"=IF({g}>0,{mid},{lo_prev})", "formula", "0.000000"); put(ws, f"{sc(pcol(t,3))}{row}", f"=IF({g}>0,{hi_prev},{mid})", "formula", "0.000000")
        put(ws, f"{sc(s_start)}{row}", 0.000001, "input"); put(ws, f"{sc(s_start+1)}{row}", 0.9999, "input")
        for t in range(1, N_ITER + 1):
            lo_prev = f"{sc(s_start if t == 1 else zcol(t-1, 2))}{row}"; hi_prev = f"{sc(s_start+1 if t == 1 else zcol(t-1, 3))}{row}"
            mid = f"{sc(zcol(t,0))}{row}"; g = f"{sc(zcol(t,1))}{row}"; Pi = f"{sc(2)}{row}"
            put(ws, mid, f"=({lo_prev}+{hi_prev})/2", "formula", "0.000000")
            put(ws, g, f"=(SUMPRODUCT(({KEYR}={sc(0)}{row})*SQRT({mid}^2+4*(1-{mid})*{IMPR}^2/{Pi}))-COUNTIF({KEYR},{sc(0)}{row})*{mid})/(2*(1-{mid}))-1", "formula", "0.000000")
            put(ws, f"{sc(zcol(t,2))}{row}", f"=IF({g}>0,{mid},{lo_prev})", "formula", "0.000000"); put(ws, f"{sc(zcol(t,3))}{row}", f"=IF({g}>0,{hi_prev},{mid})", "formula", "0.000000")
        put(ws, f"{sc(out_p)}{row}", f"=({sc(pcol(N_ITER,2))}{row}+{sc(pcol(N_ITER,3))}{row})/2", "formula", "0.000000")
        put(ws, f"{sc(out_p+1)}{row}", f"=({sc(zcol(N_ITER,2))}{row}+{sc(zcol(N_ITER,3))}{row})/2", "formula", "0.000000")
    SOLK = f"${sc(0)}${sol_first}:${sc(0)}${sol_last}"; POWK = f"${sc(out_p)}${sol_first}:${sc(out_p)}${sol_last}"; SHZ = f"${sc(out_p+1)}${sol_first}:${sc(out_p+1)}${sol_last}"
    for i in range(n["n"]):
        row = first + i; orow = f0 + i
        put(ws, f"A{row}", f"=Odds_Input!A{orow}", "link", "int"); put(ws, f"B{row}", f"=Odds_Input!B{orow}", "link"); put(ws, f"C{row}", f"=Odds_Input!C{orow}", "link")
        put(ws, f"D{row}", f"=Odds_Input!D{orow}", "link"); put(ws, f"E{row}", f"=Odds_Input!E{orow}", "link", "num1"); put(ws, f"F{row}", f"=Odds_Input!F{orow}", "link", "amer")
        put(ws, f"G{row}", f'=IF(OR(B{row}="WINS",B{row}="PLAYOFFS"),B{row}&"|"&C{row}&IF(B{row}="WINS","|"&TEXT(E{row},"0.0"),""),B{row})', "formula")
        put(ws, f"H{row}", f"=IF(F{row}>0,1+F{row}/100,1+100/ABS(F{row}))", "formula", "dec"); put(ws, f"I{row}", f"=1/H{row}", "formula", "pct2")
        put(ws, f"J{row}", f'=IF(LEFT(B{row},6)="PLAYIN",A_playin_slots,1)', "formula", "int"); put(ws, f"K{row}", f"=COUNTIF({KEYR},G{row})", "formula", "int")
        put(ws, f"L{row}", f"=SUMIF({KEYR},G{row},{IMPR})", "formula", "0.0000"); put(ws, f"M{row}", f"=L{row}/J{row}-1", "formula", "pct")
        put(ws, f"N{row}", f"=I{row}*J{row}/L{row}", "formula", "pct2")
        put(ws, f"R{row}", f"=INDEX({POWK},MATCH(G{row},{SOLK},0))", "formula", "0.000000"); put(ws, f"S{row}", f"=INDEX({SHZ},MATCH(G{row},{SOLK},0))", "formula", "0.000000")
        put(ws, f"O{row}", f"=I{row}^R{row}", "formula", "pct2")
        put(ws, f"P{row}", f"=IF(J{row}<>1,O{row},IF(L{row}<=1,I{row}/L{row},(SQRT(S{row}^2+4*(1-S{row})*I{row}^2/L{row})-S{row})/(2*(1-S{row}))))", "formula", "pct2")
        put(ws, f"Q{row}", f'=IF(A_devig="Multiplicative",N{row},IF(A_devig="Power",O{row},P{row}))', "formula", "pct2")
    REG["devig"] = dict(first=first, last=last, sol_first=sol_first, sol_last=sol_last, nkeys=len(ukeys), ukeys=ukeys)
    ws.column_dimensions.group(sc(0), sc(out_p + 1), hidden=False, outline_level=1)
    widths(ws, {"A": 8, "B": 14, "C": 28, "D": 8, "E": 8, "F": 10, "G": 34}); finish(ws, label_cols=("C", "G"))
    ws.freeze_panes = "D8"
    return ws
