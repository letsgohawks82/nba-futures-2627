"""python -m src.workbook.build : builds output/NBA_Futures_Model_v1_YYYYMMDD_MS.xlsx, recalcs in Excel (xlwings), scans errors, runs the checks."""
import datetime, json, os, shutil, sys, time
import numpy as np, pandas as pd
from openpyxl import Workbook
from .styles import *
from . import core, models, pricing_tabs as pt, consensus_tabs as ct
from .core import REG, C, ROOT

TAB_ORDER = ["Cover", "Assumptions", "Overrides", "Odds_Input", "M_DeVig", "M_TeamRatings", "M_WinDist", "M_Sims", "M_Lottery2027", "M_Players", "S_Markets", "Bet_Card", "Tier1_Card", "Exposure", "Dashboard", "Sensitivity", "Backtest",
             "Tracker", "Checks", "Data_Sources", "Raw_Teams", "Raw_Players", "Raw_Schedule", "Raw_Consensus", "Raw_MarketRef"]


def build(path):
    o = pd.read_csv(C("pricing_table.csv"))
    wb = Workbook(); wb.remove(wb.active)
    core.build_assumptions(wb); T, P = core.build_raw(wb); core.build_overrides(wb, P); odds = core.build_odds(wb); core.build_devig(wb, odds)
    ct.build_consensus(wb); models.build_team_ratings(wb); models.build_windist(wb, odds); key = models.build_sims(wb, o); models.fix_windist_links(wb); models.build_lottery(wb); models.build_players(wb)
    pt.build_smarkets(wb, o); ct.build_marketref(wb, o); pt.build_betcard(wb); pt.build_tier1(wb, o)
    Q = o[o.qualifies]
    pt.build_exposure(wb, Q); pt.build_dashboard(wb, o); pt.build_sensitivity(wb); pt.build_backtest(wb); pt.build_tracker(wb); pt.build_checks(wb, o); pt.build_sources(wb); pt.build_cover(wb, key)
    wb._sheets = [wb[n] for n in TAB_ORDER]
    for ws in wb.worksheets:           # Arial 10 default on every populated cell without an explicit font
        for row in ws.iter_rows():
            for c in row:
                if c.value is not None and c.font.name != "Arial":
                    c.font = font(BLACK)
        ws.sheet_properties.tabColor = NAVY
        for r in range(1, ws.max_row + 1):          # row heights: data 15, wrapped headers 30, spacer 6
            first = ws.cell(row=r, column=1)
            if r == 4: ws.row_dimensions[r].height = 6
            elif first.fill is not None and first.fill.fgColor is not None and first.fill.fgColor.rgb in ("001F3864", "FF1F3864", "1F3864"): ws.row_dimensions[r].height = 30
            else: ws.row_dimensions[r].height = 15
    wb.save(path)
    return path


def post_excel(path):
    """Open in Excel (xlwings): dynamic-array formulas, full recalc, error scan, checks readout."""
    import xlwings as xw
    app = xw.App(visible=False, add_book=False); app.display_alerts = False
    out = {}
    try:
        wb = app.books.open(os.path.abspath(path))
        ws = wb.sheets["Bet_Card"]; ws.range("A8").formula2 = pt.bet_card_formula()
        d = wb.sheets["Dashboard"]
        for k in range(1, 11):
            r = 16 + k
            d.range(f"B{r}").formula2 = (f'=IFERROR(INDEX(Bet_Card!$A$8#,{k},1)&" - "&INDEX(Bet_Card!$A$8#,{k},2)&" "&INDEX(Bet_Card!$A$8#,{k},3)&" "&IF(ISNUMBER(INDEX(Bet_Card!$A$8#,{k},4)),INDEX(Bet_Card!$A$8#,{k},4),""),"-")')
            d.range(f"C{r}").formula2 = f'=IFERROR(INDEX(Bet_Card!$A$8#,{k},8),0)'; d.range(f"D{r}").formula2 = f'=IFERROR(INDEX(Bet_Card!$A$8#,{k},9),0)'
        app.api.CalculateFullRebuild() if hasattr(app.api, "CalculateFullRebuild") else app.calculate()
        app.calculate()
        # number formats / widths for the spilled Bet Card
        fm = ["@", "@", "@", "0.0", "+0;-0", "0.0%", "0.0%", "0.0%", "0.0%", "@", "0.0000", "$#,##0", "0.0000", "$#,##0", "@", "@", "@"]
        fm1 = ["@", "@", "@", "0.0", "+0;-0", "0.0%", "0.0%", "0.0%", "0.0%", "0.00%", "$#,##0", "@", "@", "@"]
        n = int(wb.sheets["Dashboard"].range("B7").value or 0)
        for j, f in enumerate(fm):
            ws.range((8, 1 + j), (8 + max(n, 1) + 5, 1 + j)).number_format = f
        for rr in ws.range((8, 1), (8 + max(n, 1) + 5, 17)):
            rr.api.Font.Name = "Arial"; rr.api.Font.Size = 10
        app.calculate()
        wb.save()
        # error scan on every sheet
        errs = {}
        for sh in wb.sheets:
            ur = sh.used_range; vals = ur.options(ndim=2).value; cnt = 0; ex = []
            if vals is None: continue
            for i, row in enumerate(vals):
                for j, v in enumerate(row):
                    if (isinstance(v, str) and v.startswith("#") and v[1:4].isupper()) or (isinstance(v, int) and v in (-2146826281, -2146826246, -2146826259, -2146826288, -2146826252, -2146826265, -2146826273, -2146826239)):
                        cnt += 1
                        if len(ex) < 5: ex.append((i + ur.row, j + ur.column, v))
            errs[sh.name] = (cnt, ex)
        out["errors"] = errs
        ck = wb.sheets["Checks"]; r0 = REG["checks"]
        out["checks"] = [(ck.range(f"A{r}").value, ck.range(f"B{r}").value, ck.range(f"C{r}").value) for r in range(r0["first"], r0["last"] + 1)]
        out["overall_checks_failing"] = ck.range(f"B{r0['overall']}").value
        out["bet_count"] = wb.sheets["Dashboard"].range("B7").value; out["tier1_count"] = wb.sheets["Dashboard"].range("B6").value
        out["sim_status"] = wb.sheets["Dashboard"].range("B13").value
        # parity: Excel fair probabilities vs Python pricing table
        o = pd.read_csv(C("pricing_table.csv")); dv = REG["devig"]
        xs = wb.sheets["M_DeVig"].range(f"N{dv['first']}:Q{dv['last']}").value
        xs = np.array(xs, float)
        out["devig_parity_max_abs"] = {"mult": float(np.nanmax(np.abs(xs[:, 0] - o.fair_mult.values))), "power": float(np.nanmax(np.abs(xs[:, 1] - o.fair_power.values))), "shin": float(np.nanmax(np.abs(xs[:, 2] - o.fair_shin.values))), "selected": float(np.nanmax(np.abs(xs[:, 3] - o.fair.values)))}
        tr = REG["tr"]; tw = np.array(wb.sheets["M_TeamRatings"].range(f"K{tr['first']}:K{tr['last']}").value, float); tp = np.array(wb.sheets["M_TeamRatings"].range(f"S{tr['first']}:S{tr['last']}").value, float)
        out["team_wins_parity_max_abs"] = float(np.max(np.abs(tw - tp))); out["wins_sum"] = float(tw.sum())
        # S_Markets EV parity vs Python for modelled rows
        sm = REG["smk"]; rows = sm["rowmap"]; sh = wb.sheets["S_Markets"]
        ev_x = []; ev_p = []
        for idx, rr in rows.items():
            if bool(o.loc[idx, "modelled"]):
                v = sh.range(f"L{rr}").value; ev_x.append(v); ev_p.append(o.loc[idx, "ev"] if not (o.loc[idx, "market"] == "PLAYOFFS") else o.loc[idx, "ev"])
        out["ev_parity_max_abs"] = float(np.nanmax(np.abs(np.array(ev_x, float) - np.array(ev_p, float))))
        qx = int(sum(1 for idx, rr in rows.items() if sh.range(f"W{rr}").value is True)); out["qualifying_excel"] = qx; out["qualifying_python"] = int(o.qualifies.sum())
        wb.save(); wb.close()
    finally:
        app.quit()
    return out


if __name__ == "__main__":
    stamp = datetime.date.today().strftime("%Y%m%d")
    path = os.path.join(ROOT, "output", f"NBA_Futures_Model_v1_{stamp}_MS.xlsx")
    os.makedirs(os.path.dirname(path), exist_ok=True)
    if os.path.exists(path):          # retain prior versions
        shutil.copy(path, path.replace(".xlsx", f"_prev{datetime.datetime.now().strftime('%H%M%S')}.xlsx"))
    t = time.time(); build(path); print("built", path, round(time.time() - t, 1), "s")
    res = post_excel(path); print(json.dumps(res, indent=1, default=str)[:6000])
    json.dump(res, open(C("workbook_build_report.json"), "w"), indent=1, default=str)
