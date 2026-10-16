"""Six-bug audit + formatting audit of the built workbook. python -m src.workbook.audit [path]"""
import glob, json, os, re, sys
from collections import Counter, defaultdict
import numpy as np, openpyxl
from openpyxl.utils import get_column_letter as CL, column_index_from_string as CI

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
REF = re.compile(r"(\$?)([A-Z]{1,3})(\$?)(\d+)")
NUM = re.compile(r"(?<![A-Za-z_$\d.])(\d+\.?\d*)(?![\d.]*[A-Za-z_(])")
ALLOWED = {"0", "1", "2", "3", "4", "5", "6", "7", "8", "10", "100", "0.5", "0.000001", "0.0001", "48", "-1"}     # structural: 0/1 flags, 100 (American odds), 48 (min/game), pairs, tolerances


def relative(formula, row, col):
    def sub(m):
        cabs, c, rabs, r = m.groups()
        cc = c if cabs else f"C[{CI(c) - col}]"
        rr = r if rabs else f"R[{int(r) - row}]"
        return f"{cc}{rr}"
    return REF.sub(sub, formula)


def run(path):
    wb = openpyxl.load_workbook(path)           # formulas
    out = {"hardcodes": {}, "inconsistent": {}, "format": {}}
    for ws in wb.worksheets:
        # (1) hardcoded numeric literals in formulas
        lit = Counter(); ex = {}
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if isinstance(v, str) and v.startswith("="):
                    body = re.sub(r'"[^"]*"', '""', v)
                    for m in NUM.finditer(body):
                        t = m.group(1)
                        if t not in ALLOWED:
                            lit[t] += 1; ex.setdefault(t, c.coordinate)
        out["hardcodes"][ws.title] = {k: (n, ex[k]) for k, n in lit.most_common(8)}
        # (2) inconsistent formulas down columns (relative-form patterns per column; blocks are expected to differ only at block boundaries)
        incons = {}
        cols = defaultdict(list)
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    cols[c.column].append((c.row, relative(c.value, c.row, c.column)))
        for col, items in cols.items():
            pats = Counter(p for _, p in items)
            if len(items) > 12 and len(pats) > 1:
                top, n = pats.most_common(1)[0]
                minor = sum(v for p, v in pats.items() if p != top)
                if minor / len(items) > 0.02:
                    incons[CL(col)] = dict(patterns=len(pats), dominant_share=round(n / len(items), 3), n=len(items))
        out["inconsistent"][ws.title] = incons
        # formatting
        bad_font = sum(1 for row in ws.iter_rows() for c in row if c.value is not None and (c.font.name != "Arial" or c.font.size != 10) and c.font.size not in (12, 16))
        merged = len(ws.merged_cells.ranges)
        out["format"][ws.title] = dict(non_arial10_cells=bad_font, merged=merged, gridlines=ws.sheet_view.showGridLines,
                                       min_label_width=min([ws.column_dimensions[CL(i)].width or 0 for i in range(1, 2)]), tab_header_rows=bool(ws["A1"].value and ws["A2"].value and ws["A3"].value))
    # (3)/(4) circular refs and external links via Excel
    import xlwings as xw
    app = xw.App(visible=False, add_book=False); app.display_alerts = False
    try:
        xb = app.books.open(os.path.abspath(path))
        circ = {}
        for sh in xb.sheets:
            try:
                r = sh.api.CircularReference
                circ[sh.name] = r.Address if r is not None else None
            except Exception:
                circ[sh.name] = None
        out["circular"] = {k: v for k, v in circ.items() if v}
        ls = xb.api.LinkSources(1)
        out["external_links"] = list(ls) if ls else []
        xb.close()
    finally:
        app.quit()
    # (5)/(6) static convention checks recorded by name
    out["sign_conventions"] = ["EV = p x decimal - 1 (positive = value)", "American odds: positive = profit on $100, negative = stake to win $100 (converted by decimal formula)", "Edge = model - fair (positive = model above market)",
                               "CLV = closing implied - placed implied (positive = beat the close)", "Kelly floored at 0; stakes never negative", "Wins residual = actual - predicted"]
    # (5b) numeric sign checks on cached values; (6) summary tabs linked (formula cells that reference other tabs)
    vb = openpyxl.load_workbook(path, data_only=True); mp = vb["M_Players"]; bad_ev = bad_stake = 0; n = 0
    for r in mp.iter_rows(min_row=8, values_only=True):
        if r[0] is None or not isinstance(r[4], (int, float)): continue
        n += 1; bad_ev += abs((r[6] * r[4] - 1) - r[9]) > 1e-9; bad_stake += (isinstance(r[15], (int, float)) and r[15] < 0)
    sm = vb["S_Markets"]; neg = sum(1 for row in sm.iter_rows(min_row=8, min_col=35, max_col=36, values_only=True) for v in row if isinstance(v, (int, float)) and v < 0)
    out["sign_checks"] = dict(m_players_rows=n, ev_formula_mismatch=int(bad_ev), negative_player_stakes=int(bad_stake), negative_team_stakes=int(neg))
    linked = {}
    for sh, cells in (("Dashboard", ["B6", "B7", "B8", "B9", "B10", "B11", "B12", "B13", "B17"]), ("Exposure", ["B14", "B15"]), ("Tier1_Card", ["A8", "A41"]), ("Bet_Card", ["A8"])):
        for c in cells:
            v = wb[sh][c].value; v = getattr(v, "text", v)
            linked[f"{sh}!{c}"] = bool(isinstance(v, str) and v.startswith("=") and re.search(r"[A-Za-z_]+!", v))
    out["summary_links"] = linked
    out["summary_hardcoded_engine_outputs"] = ["Exposure!B8:B13 (joint-Kelly engine summary from the sims)", "M_Sims (engine), M_Players probabilities/stakes (engine)"]
    return out


if __name__ == "__main__":
    path = sys.argv[1] if len(sys.argv) > 1 else sorted(glob.glob(os.path.join(ROOT, "output", "NBA_Futures_Model_v1_*_MS.xlsx")))[-1]
    r = run(path)
    json.dump(r, open(os.path.join(ROOT, "data/clean/workbook_audit.json"), "w"), indent=1)
    print("HARDCODES:"); [print(" ", s, v) for s, v in r["hardcodes"].items() if v]
    print("INCONSISTENT COLUMNS:"); [print(" ", s, v) for s, v in r["inconsistent"].items() if v]
    print("SIGN:", r["sign_checks"]); print("SUMMARY LINKS:", {k: v for k, v in r["summary_links"].items() if not v} or "all linked")
    print("FORMAT:"); [print(" ", s, v) for s, v in r["format"].items() if v["non_arial10_cells"] or v["merged"] or not v["tab_header_rows"] or v["gridlines"]]
    print("CIRCULAR:", r["circular"], "| EXTERNAL LINKS:", r["external_links"])
