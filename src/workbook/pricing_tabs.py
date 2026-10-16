"""Workbook pricing/reporting tabs: S_Markets, Bet_Card (formulas set via Excel), Exposure, Dashboard, Sensitivity, Backtest, Tracker, Checks, Data_Sources, Cover."""
import json, math, os
import numpy as np, pandas as pd
from openpyxl.utils import get_column_letter as CL
from openpyxl.formatting.rule import CellIsRule, FormulaRule, ColorScaleRule
from openpyxl.styles import Font, PatternFill
from openpyxl.chart import BarChart, ScatterChart, Reference, Series
from openpyxl.worksheet.datavalidation import DataValidation
from .styles import *
from .core import REG, C, ROOT
from ..data.teams import NAME2ABBR

MARKET_ORDER = ["TITLE", "CONF_EAST", "CONF_WEST", "DIV_ATLANTIC", "DIV_CENTRAL", "DIV_SOUTHEAST", "DIV_NORTHWEST", "DIV_PACIFIC", "DIV_SOUTHWEST", "CUP", "CUP_EAST_A", "CUP_EAST_B", "CUP_EAST_C", "CUP_WEST_A", "CUP_WEST_B", "CUP_WEST_C",
                "PLAYIN_EAST", "PLAYIN_WEST", "PLAYOFFS", "WINS", "MVP", "DPOY", "ROY", "LDR_PPG", "LDR_RPG", "LDR_APG", "LDR_SPG", "LDR_BPG", "LDR_TPG", "MIP", "SMOY", "COY", "CLUTCH", "CUP_MVP"]
GREEN_FILL = PatternFill("solid", bgColor="C6E0B4", fgColor="C6E0B4"); RED_FILL = PatternFill("solid", bgColor="F4B6B6", fgColor="F4B6B6")


def build_smarkets(wb, o):
    ws = wb.create_sheet("S_Markets"); sheet_header(ws, "S_Markets - one block per market: book price, fair probability, model probability, EV, Kelly, edge CI", span=12)
    put(ws, "A5", "Edge colours: green = EV >= min EV with CI low > 0; red = EV < 0. Imputed (orange) sides are never bet. Model probabilities are sim-based (M_Sims); WINS also shown live in M_WinDist.", "text", italic=True)
    put(ws, "L5", "=Sim_Stale", "link"); ws["L5"].font = font(RED, bold=True)
    hdr = ["Market", "Selection", "Team", "Side", "Line", "Book American", "Decimal", "Flag", "Fair prob (de-vig)", "Model prob", "Model / fair price (American)", "EV per $1", "Edge (model-fair, pp)", "EV CI low (80%)", "EV CI high (80%)", "P(EV>0)",
           "Staking prob (shrunk)", "Naive Kelly (frac)", "Naive stake $", "Portfolio Kelly (frac)", "Portfolio stake $", "Robust OK", "Qualifies", "Lock (CST)", "Rule", "Row id", "Sim status", "Live analytic P (WINS, reprices with Overrides)", "Live analytic EV (WINS)", "Zero-sum model P (WINS, means rescaled to book sum)", "Zero-sum EV (WINS)", "EV at shrunk probability", "EV CI low at shrunk probability", "Tier 1 (live)", "Tier 1 stake (frac)", "Tier 1 stake $", "Named driver", "Portfolio Kelly (numeric, for sorting)", "Tier"]
    headers(ws, 7, hdr, height=30)
    S = REG["sims_team"]; cm = S["cols"]; sf, sl = S["first"], S["last"]
    pm = REG["sims_pmf"]; rk0 = REG["sims_rank_col0"]; ci = REG["sims_ci"]; od = REG["odds"]; dv = REG["devig"]
    teamcol = lambda key: f"INDEX(M_Sims!${cm[key]}${sf}:${cm[key]}${sl},MATCH($C{{r}},M_Sims!$A${sf}:$A${sl},0))"
    r = 8; rowmap = {}; first_data = 8
    for m in MARKET_ORDER:
        sub = o[o.market == m]
        if sub.empty: continue
        put(ws, f"A{r}", m, "text", bold=True); put(ws, f"B{r}", f"{len(sub)} selections - {sub.rule_text.iloc[0]}", "text", italic=True)
        for col in range(1, 40): ws.cell(row=r, column=col).border = Border(bottom=THIN)
        r += 1
        for idx, x in sub.iterrows():
            rowmap[idx] = r; orow = od["first"] + idx; dvr = dv["first"] + idx; cir = ci["first"] + idx
            modelled = bool(x.modelled)
            put(ws, f"A{r}", m, "text"); put(ws, f"B{r}", f"=Odds_Input!C{orow}", "link"); put(ws, f"C{r}", x.team if isinstance(x.team, str) else "-", "text"); put(ws, f"D{r}", f"=Odds_Input!D{orow}", "link")
            put(ws, f"E{r}", f"=Odds_Input!E{orow}", "link", "num1"); put(ws, f"F{r}", f"=Odds_Input!F{orow}", "link", "amer"); put(ws, f"G{r}", f"=M_DeVig!H{dvr}", "link", "dec"); put(ws, f"H{r}", f"=Odds_Input!G{orow}", "link")
            put(ws, f"I{r}", f"=M_DeVig!Q{dvr}", "link", "pct")
            if modelled:
                t = lambda k: teamcol(k).format(r=r)
                if m == "TITLE": f = t("p_title")
                elif m in ("CONF_EAST", "CONF_WEST"): f = t("p_conference")
                elif m.startswith("DIV_"): f = t("p_division")
                elif m == "CUP": f = t("p_cup")
                elif m.startswith("CUP_"): f = t("p_cup_group")
                elif m.startswith("PLAYIN"): f = f"SUM(INDEX(M_Sims!${CL(rk0+7)}${sf}:${CL(rk0+10)}${sl},MATCH($C{r},M_Sims!$A${sf}:$A${sl},0),0))"
                elif m == "PLAYOFFS":
                    q = f'IF(A_podef="Bracket",{t("p_bracket")},{t("p_playoffs")})'; f = q if x.side == "YES" else f"1-({q})"
                else:   # WINS
                    hdr_r = f"M_Sims!${CL(pm['c0'])}${pm['hdr']}:${CL(pm['c1'])}${pm['hdr']}"
                    over = f"SUMPRODUCT(({hdr_r}>E{r})*INDEX(M_Sims!${CL(pm['c0'])}${pm['first']}:${CL(pm['c1'])}${pm['last']},MATCH($C{r},M_Sims!$A${pm['first']}:$A${pm['last']},0),0))"
                    f = over if x.side == "OVER" else f"1-{over}"
                put(ws, f"J{r}", "=" + f, "link", "pct")
                put(ws, f"K{r}", f"=IF(J{r}>=0.5,-100*J{r}/(1-J{r}),100*(1-J{r})/J{r})", "formula", "amer")
                put(ws, f"L{r}", f"=J{r}*G{r}-1", "formula", "pct"); put(ws, f"M{r}", f"=(J{r}-I{r})*100", "formula", "num1")
                if m == "PLAYOFFS":
                    put(ws, f"N{r}", f'=IF(A_podef="Bracket",M_Sims!B{cir},M_Sims!F{cir})', "link", "pct"); put(ws, f"O{r}", f'=IF(A_podef="Bracket",M_Sims!C{cir},M_Sims!G{cir})', "link", "pct")
                else:
                    put(ws, f"N{r}", f"=M_Sims!B{cir}", "link", "pct"); put(ws, f"O{r}", f"=M_Sims!C{cir}", "link", "pct")
                put(ws, f"P{r}", f"=M_Sims!D{cir}", "link", "pct")
                put(ws, f"Q{r}", f"=(1-A_shrink)*J{r}+A_shrink*I{r}", "formula", "pct"); put(ws, f"R{r}", f'=IF(A_tier2_on=1,MAX(0,(Q{r}*G{r}-1)/(G{r}-1))*A_kelly,"-")', "formula", "0.0000")
                put(ws, f"S{r}", f'=IF(OR(A_bankroll="",A_tier2_on<>1),"-",MIN(A_maxbet,MAX(0,(Q{r}*G{r}-1)/(G{r}-1))*A_kelly*A_bankroll))', "formula", "money")
                put(ws, f"T{r}", f'=IF(A_tier2_on=1,M_Sims!M{cir},"-")', "link", "0.0000"); put(ws, f"U{r}", f'=IF(OR(A_bankroll="",A_tier2_on<>1),"-",MIN(A_maxbet,M_Sims!M{cir}*A_bankroll))', "formula", "money")
                put(ws, f"V{r}", f"=M_Sims!H{cir}", "link", "int")
                put(ws, f"AF{r}", f"=Q{r}*G{r}-1", "formula", "pct"); put(ws, f"AG{r}", f"=(1-A_shrink)*N{r}+A_shrink*(I{r}*G{r}-1)", "formula", "pct")
                if m == "WINS":
                    rowr = f"INDEX(M_Sims!${CL(pm['c0'])}${pm['first']}:${CL(pm['c1'])}${pm['last']},MATCH($C{r},M_Sims!$A${pm['first']}:$A${pm['last']},0),0)"
                    hdr_r2 = f"M_Sims!${CL(pm['c0'])}${pm['hdr']}:${CL(pm['c1'])}${pm['hdr']}"
                    ov_zs = f"(SUMPRODUCT(({hdr_r2}>=INT(E{r})+1)*{rowr})+A_zs_delta*INDEX({rowr},1,INT(E{r})+2))"
                    put(ws, f"AD{r}", "=" + (ov_zs if x.side == "OVER" else f"1-{ov_zs}"), "link", "pct"); put(ws, f"AE{r}", f"=AD{r}*G{r}-1", "formula", "pct")
                else:
                    put(ws, f"AD{r}", "-", "text", align="right"); put(ws, f"AE{r}", "-", "text", align="right")
                zs = f'IF(AND(A{r}="WINS",A_zs_on=1),AND(AE{r}>=A_minev,M_Sims!O{cir}>0),TRUE)'
                rt_ = REG["raw_teams"]
                cont = f"INDEX(Raw_Teams!$F${rt_['first']}:$F${rt_['last']},MATCH($C{r},Raw_Teams!$A${rt_['first']}:$A${rt_['last']},0))"
                longshot = f"IF(I{r}<A_longthr,(1-A_shrink)*M_Sims!K{cir}+A_shrink*(I{r}*G{r}-1)>=A_minev,TRUE)"
                contf = f'IF(AND(OR(A{r}="WINS",A{r}="PLAYOFFS",LEFT(A{r},4)="DIV_"),{cont}<A_contthr),AND(((1-A_contshrink)*J{r}+A_contshrink*I{r})*G{r}-1>=A_minev,(1-A_contshrink)*N{r}+A_contshrink*(I{r}*G{r}-1)>0),TRUE)'
                put(ws, f"W{r}", f'=AND(H{r}<>"IMPUTED_20C_LADDER",L{r}>=A_minev,N{r}>0,V{r}=1,IF(A{r}="PLAYOFFS",A_poconf=1,TRUE),{zs},AF{r}>=A_minev,AG{r}>0,{longshot},{contf},M_Sims!P{cir}=0)', "formula")
                put(ws, f"AH{r}", f"=AND(W{r},I{r}>=A_t1_fair,((1-A_t1_shrink)*J{r}+A_t1_shrink*I{r})*G{r}-1>=A_minev,(1-A_t1_shrink)*N{r}+A_t1_shrink*(I{r}*G{r}-1)>0,M_Sims!Q{cir}=1)", "formula")
                put(ws, f"AI{r}", f"=IF(AH{r},M_Sims!R{cir}*A_t1_cap/A_t1_engcap,0)", "formula", "0.0000")
                put(ws, f"AJ{r}", f'=IF(OR(A_bankroll="",NOT(AH{r})),"-",MIN(A_maxbet,AI{r}*A_bankroll))', "formula", "money")
                put(ws, f"AK{r}", f'=IF(AH{r},M_Sims!S{cir},"-")', "link")
                put(ws, f"AL{r}", f"=M_Sims!M{cir}", "link", "0.0000"); put(ws, f"AM{r}", f'=IF(W{r},IF(AH{r},"Tier 1","Tier 2"),"-")', "formula")
            else:
                put(ws, f"J{r}", "-", "text", align="right"); put(ws, f"K{r}", f"=IF(I{r}>=0.5,-100*I{r}/(1-I{r}),100*(1-I{r})/I{r})", "formula", "amer")
                for col in "LMNOPQRSTUV": put(ws, f"{col}{r}", "-", "text", align="right")
                put(ws, f"W{r}", False, "text")
                for col in ("AD", "AE", "AF", "AG", "AI", "AJ", "AK", "AM"): put(ws, f"{col}{r}", "-", "text", align="right")
                put(ws, f"AH{r}", False, "text"); put(ws, f"AL{r}", 0, "text")
            if m == "WINS":
                wd = REG["windist"]
                put(ws, f"AB{r}", f'=INDEX(IF(D{r}="OVER",M_WinDist!$E${wd["first"]}:$E${wd["last"]},M_WinDist!$F${wd["first"]}:$F${wd["last"]}),MATCH($C{r},M_WinDist!$A${wd["first"]}:$A${wd["last"]},0))', "link", "pct")
                put(ws, f"AC{r}", f"=AB{r}*G{r}-1", "formula", "pct")
            else:
                put(ws, f"AB{r}", "-", "text", align="right"); put(ws, f"AC{r}", "-", "text", align="right")
            put(ws, f"X{r}", x.lock_cst, "text"); put(ws, f"Y{r}", x.rule_text, "text"); put(ws, f"Z{r}", int(idx), "text"); put(ws, f"AA{r}", "=Sim_Stale", "link")
            if isinstance(x.flag, str):
                for col in "ABCDEFGH": ws[f"{col}{r}"].fill = PatternFill("solid", fgColor=ORANGE)
            r += 1
        r += 1
    last = r
    REG["smk"] = dict(first=8, last=last, rowmap=rowmap)
    ws.conditional_formatting.add(f"L8:L{last}", FormulaRule(formula=['AND(ISNUMBER(L8),L8>=A_minev,ISNUMBER(N8),N8>0)'], fill=GREEN_FILL))
    ws.conditional_formatting.add(f"L8:L{last}", FormulaRule(formula=['AND(ISNUMBER(L8),L8<0)'], fill=RED_FILL))
    ws.conditional_formatting.add(f"AA8:AA{last}", FormulaRule(formula=['LEFT(AA8,5)="STALE"'], font=Font(name="Arial", size=10, color=RED, bold=True)))
    widths(ws, {"A": 16, "B": 28, "C": 8, "D": 8, "X": 24, "Y": 60, "AA": 20}); finish(ws); ws.freeze_panes = "C8"
    ws.column_dimensions.group("Z", "Z", hidden=False)
    return ws


def build_betcard(wb):
    ws = wb.create_sheet("Bet_Card"); sheet_header(ws, "Bet_Card - all qualifying bets (dynamic FILTER/SORTBY over S_Markets)", span=15)
    sm = REG["smk"]
    put(ws, "A5", "Tier 1 = conservative subset (see Tier1_Card, has stakes). Tier 2 = everything on this card, NO stake by default (Assumptions: tier2_staking_on). Qualifies = EV >= min EV AND edge CI low > 0 AND robustness filter AND not imputed AND (WINS) survives the zero-sum rescale AND the edge survives the shrink toward the market. Stake = portfolio Kelly (correlation-aware), naive shown beside it. Lock times in CST.", "text", italic=True)
    put(ws, "A6", "=Sim_Stale", "link"); ws["A6"].font = font(RED, bold=True)
    labels = ["Market", "Selection", "Side", "Line", "Book American", "Fair prob", "Model prob", "EV per $1", "EV CI low", "Tier", "Naive Kelly", "Naive stake $", "Portfolio Kelly", "Portfolio stake $", "Lock (CST)", "Rule"]
    headers(ws, 7, labels)
    REG["bc"] = dict(first=8, cols=labels)
    widths(ws, {"A": 16, "B": 28, "N": 24, "O": 60}); finish(ws); ws.freeze_panes = "C8"
    return ws


def bet_card_formula():
    sm = REG["smk"]
    d = f"S_Markets!$A${sm['first']}:$AM${sm['last']}"
    cols = [1, 2, 4, 5, 6, 9, 10, 12, 14, 39, 18, 19, 20, 21, 24, 25]
    pick = ",".join(f"INDEX(s,0,{c})" for c in cols)
    return (f'=LET(d,{d},q,INDEX(d,0,23),f,FILTER(d,q=TRUE,"No qualifying bets"),'
            f'IF(COLUMNS(f)<2,f,LET(s,SORTBY(f,INDEX(f,0,38),-1),HSTACK({pick}))))')


def tier1_formula():
    sm = REG["smk"]
    d = f"S_Markets!$A${sm['first']}:$AM${sm['last']}"
    cols = [1, 2, 4, 5, 6, 9, 10, 12, 14, 35, 36, 37, 24, 25]
    pick = ",".join(f"INDEX(s,0,{c})" for c in cols)
    return (f'=LET(d,{d},q,INDEX(d,0,34),f,FILTER(d,q=TRUE,"No Tier 1 bets"),'
            f'IF(COLUMNS(f)<2,f,LET(s,SORTBY(f,INDEX(f,0,35),-1),HSTACK({pick}))))')


def player_tier1_formula():
    mp = REG["m_players"]; d = f"M_Players!$A${mp['first']}:$U${mp['last']}"
    pick = ",".join(f"INDEX(s,0,{c})" for c in (1, 2, 3, 4, 6, 7, 10, 12, 16, 17, 20))
    return (f'=LET(d,{d},f,FILTER(d,INDEX(d,0,15)=TRUE,"No player Tier 1 bets"),'
            f'IF(COLUMNS(f)<2,f,LET(s,SORTBY(f,INDEX(f,0,16),-1),HSTACK({pick}))))')


def build_tier1(wb):
    ws = wb.create_sheet("Tier1_Card"); sheet_header(ws, "Tier1_Card - conservative subset of the Bet Card (stakes shown here by default)", span=14)
    put(ws, "A5", "Tier 1 filters (all must hold): (1) on the Bet Card; (2) edge survives shrink 0.75 (model weight 0.25); (3) market fair probability >= 5%; (4) not against a market top-5 team or for a bottom-5 team (wins, playoffs, division, conference, title); (5) a NAMED DRIVER - verifiable arrival / departure / availability facts of >= 1.0 SRS covering >= 60% of the model-vs-market gap; (6) total exposure capped at 10% of bankroll.", "text", italic=True)
    put(ws, "A6", "=Sim_Stale", "link"); ws["A6"].font = font(RED, bold=True)
    labels = ["Market", "Selection", "Side", "Line", "Book American", "Fair prob", "Model prob", "EV per $1", "EV CI low", "Tier 1 stake (frac of bankroll)", "Tier 1 stake $", "Named driver (player / roster facts)", "Lock (CST)", "Rule"]
    headers(ws, 7, labels)
    put(ws, "A39", "Player Tier 1 (M_Players): Tier 2 test + fair >= A_t1_fair + edge survives A_t1_shrink + holds at both bounds of the stat-leader sharpening exponent CI. Stakes are the JOINT portfolio Kelly with the team bets above (shared sims, total capped at A_t1_cap).", "text", italic=True)
    headers(ws, 40, ["Market", "Selection", "Team", "Book American", "Fair prob", "Model prob", "EV per $1", "EV low (shrunk)", "Tier 1 stake (frac of bankroll)", "Tier 1 stake $", "Driver (projection facts)"])
    widths(ws, {"A": 16, "B": 28, "L": 90, "M": 24, "N": 60}); finish(ws); ws.freeze_panes = "C8"
    return ws


def build_exposure(wb, Q):
    ws = wb.create_sheet("Exposure"); sheet_header(ws, "Exposure - portfolio P&L distribution from the sims (per $1 of bankroll; x bankroll when entered)", span=8)
    t1 = json.load(open(C("joint_summary.json"))); s = json.load(open(C("portfolio_summary.json")))["summary"]
    put(ws, "A5", "Engine output. Tier 1 (team + player bets) is the staked card: portfolio Kelly solved JOINTLY on the shared season sims, cap 10% of bankroll. Tier 2 (full team Bet Card) is shown only as 'if staked'. Rerun `run.py price` after changing thresholds.", "text", italic=True)
    headers(ws, 7, ["Metric", "Per $1 bankroll", "In $ (needs bankroll)"])
    t1_items = [("TIER 1 - bets (team + player)", t1["n_bets"], "int"), ("Tier 1 - total stake, joint Kelly (fraction of bankroll)", t1["total_stake_frac"], "pct2"), ("Tier 1 - expected return (model world)", t1["exp_return_model"], "pct2"),
                ("Tier 1 - P(profit)", t1["p_profit"], "pct"), ("Tier 1 - 5th percentile", t1["p05"], "pct2"), ("Tier 1 - worst simulated season", t1["worst"], "pct2")]
    for i, (k, v, f) in enumerate(t1_items):
        r = 8 + i; put(ws, f"A{r}", k, "text"); put(ws, f"B{r}", float(v), "input", f)
        if f == "pct2": put(ws, f"C{r}", f'=IF(A_bankroll="","-",B{r}*A_bankroll)', "formula", "money")
        else: put(ws, f"C{r}", "-", "text", align="right")
    mp_ = REG["m_players"]; smk_ = REG["smk"]
    put(ws, "A14", "Tier 1 stake - team bets (live)", "text"); put(ws, "B14", f'=SUMIFS(S_Markets!$AI${smk_["first"]}:$AI${smk_["last"]},S_Markets!$AH${smk_["first"]}:$AH${smk_["last"]},TRUE)', "link", "pct2"); put(ws, "C14", '=IF(A_bankroll="","-",B14*A_bankroll)', "formula", "money")
    put(ws, "A15", "Tier 1 stake - player bets (live)", "text"); put(ws, "B15", f'=SUMIFS(M_Players!$P${mp_["first"]}:$P${mp_["last"]},M_Players!$O${mp_["first"]}:$O${mp_["last"]},TRUE)', "link", "pct2"); put(ws, "C15", '=IF(A_bankroll="","-",B15*A_bankroll)', "formula", "money")
    t2 = [("TIER 2 (full card, only if staked) - bets", s["n_bets"], "int"), ("Tier 2 - total stake if staked", s["total_stake_frac"], "pct"), ("Tier 2 - expected return (model world)", s["exp_return_frac"], "pct"), ("Tier 2 - expected return at shrunk probabilities", s["exp_return_tilted"], "pct"),
          ("Tier 2 - P(profit)", s["p_profit"], "pct"), ("Tier 2 - 5th percentile", s["p05"], "pct"), ("Tier 2 - 1st percentile", s["p01"], "pct"), ("Tier 2 - worst simulated season", s["worst"], "pct")]
    for i, (k, v, f) in enumerate(t2):
        r = 16 + i; put(ws, f"A{r}", k, "text"); put(ws, f"B{r}", float(v), "input", f); put(ws, f"C{r}", "-", "text", align="right")
    put(ws, "A26", "Tier 1 exposure by team (fraction of bankroll)", "text", bold=True); headers(ws, 27, ["Team", "Tier 1 stake (frac)", "Stake $"])
    sm = REG["smk"]
    for i, t in enumerate(sorted(set(NAME2ABBR.values()))):
        r = 28 + i; put(ws, f"A{r}", t, "text")
        put(ws, f"B{r}", f'=SUMIFS(S_Markets!$AI${sm["first"]}:$AI${sm["last"]},S_Markets!$C${sm["first"]}:$C${sm["last"]},A{r},S_Markets!$AH${sm["first"]}:$AH${sm["last"]},TRUE)', "link", "pct2")
        put(ws, f"C{r}", f'=IF(A_bankroll="","-",B{r}*A_bankroll)', "formula", "money")
    put(ws, "A58", "Total", "text", bold=True); put(ws, "B58", "=SUM(B28:B57)", "formula", "pct2"); total_border(ws, 58, 1, 3)
    REG["exp_team"] = (28, 57)
    put(ws, "E7", "Correlation of Tier 2 bet outcomes (sim-based)", "text", bold=True)
    if Q is not None and len(Q) > 1 and os.path.exists(C("portfolio_corr.csv")):
        Cm = pd.read_csv(C("portfolio_corr.csv"), index_col=0); n = min(len(Cm), 40)
        for j in range(n):
            put(ws, f"{CL(6+j)}8", Cm.columns[j][:26], "header")
            put(ws, f"E{9+j}", Cm.index[j][:40], "text")
            for k in range(n):
                put(ws, f"{CL(6+k)}{9+j}", float(Cm.iloc[j, k]), "input", "0.00")
        ws.conditional_formatting.add(f"F9:{CL(5+n)}{8+n}", ColorScaleRule(start_type="num", start_value=-0.3, start_color="FFFFFF", end_type="num", end_value=1, end_color="8EA9DB"))
    jc = pd.read_csv(C("joint_corr.csv"), index_col=0); n1 = len(jc); r0 = 54
    put(ws, f"E{r0-1}", "Correlation of the Tier 1 bet outcomes (team + player, shared sims)", "text", bold=True)
    for j in range(n1):
        put(ws, f"{CL(6+j)}{r0}", jc.columns[j][:26], "header"); put(ws, f"E{r0+1+j}", jc.index[j][:40], "text")
        for k in range(n1): put(ws, f"{CL(6+k)}{r0+1+j}", float(jc.iloc[j, k]), "input", "0.00")
    widths(ws, {"A": 46, "B": 18, "C": 20, "E": 42}); finish(ws)


def build_dashboard(wb, o):
    ws = wb.create_sheet("Dashboard"); sheet_header(ws, "Dashboard - headline metrics", span=12)
    sm = REG["smk"]; ps = json.load(open(C("portfolio_summary.json")))["summary"]
    W = f"S_Markets!$W${sm['first']}:$W${sm['last']}"; T = f"S_Markets!$T${sm['first']}:$T${sm['last']}"; R = f"S_Markets!$R${sm['first']}:$R${sm['last']}"
    headers(ws, 5, ["Metric", "Value"])
    AH = f"S_Markets!$AH${sm['first']}:$AH${sm['last']}"; AI = f"S_Markets!$AI${sm['first']}:$AI${sm['last']}"
    MP = REG["m_players"]; PO = f"M_Players!$O${MP['first']}:$O${MP['last']}"; PP = f"M_Players!$P${MP['first']}:$P${MP['last']}"
    mets = [("Tier 1 bet count (team + player)", f"=COUNTIF({AH},TRUE)+COUNTIF({PO},TRUE)", "int"), ("Tier 2 (full card) bet count - no stake by default", f"=COUNTIF({W},TRUE)", "int"), ("Tier 1 total stake (fraction of bankroll)", f"=SUMIFS({AI},{AH},TRUE)+SUMIFS({PP},{PO},TRUE)", "pct2"),
            ("Tier 1 total stake $", f'=IF(A_bankroll="","-",(SUMIFS({AI},{AH},TRUE)+SUMIFS({PP},{PO},TRUE))*A_bankroll)', "money"), ("Tier 1 expected return (model world)", "=Exposure!B10", "pct2"), ("Tier 1 P(profit)", "=Exposure!B11", "pct"),
            ("Largest single-team Tier 1 exposure (frac)", f"=MAX(Exposure!B{REG['exp_team'][0]}:B{REG['exp_team'][1]})", "pct2"), ("Sim status", "=Sim_Stale", None)]
    for i, (k, f, fm) in enumerate(mets):
        put(ws, f"A{6+i}", k, "text"); put(ws, f"B{6+i}", f, "link", fm)
    put(ws, "A15", "Top-10 Tier 2 bets (ranked by portfolio Kelly)", "text", bold=True); headers(ws, 16, ["#", "Bet", "EV per $1", "EV CI low"])
    for k in range(1, 11):
        r = 16 + k
        put(ws, f"A{r}", k, "text"); put(ws, f"B{r}", "-", "link"); put(ws, f"C{r}", 0, "link", "pct"); put(ws, f"D{r}", 0, "link", "pct")      # replaced by dynamic-array formulas in Excel post-processing
    ch = BarChart(); ch.type = "bar"; ch.title = "Top-10 bets: EV per $1"; ch.height = 8; ch.width = 16
    ch.add_data(Reference(ws, min_col=3, min_row=16, max_row=26), titles_from_data=True); ch.set_categories(Reference(ws, min_col=2, min_row=17, max_row=26)); ws.add_chart(ch, "F5")
    # wins scatter helper
    wd = REG["windist"]
    put(ws, "A30", "Model vs market wins (M_WinDist)", "text", bold=True); headers(ws, 31, ["Team", "Book line", "Model mean wins"])
    for i in range(wd["n"]):
        r = 32 + i; put(ws, f"A{r}", f"=M_WinDist!A{wd['first']+i}", "link"); put(ws, f"B{r}", f"=M_WinDist!B{wd['first']+i}", "link", "num1"); put(ws, f"C{r}", f"=M_WinDist!C{wd['first']+i}", "link", "num1")
    sc = ScatterChart(); sc.title = "Model wins vs book win totals"; sc.style = 13; sc.height = 8; sc.width = 16; sc.x_axis.title = "Book line"; sc.y_axis.title = "Model mean wins"
    s = Series(Reference(ws, min_col=3, min_row=32, max_row=31 + wd["n"]), Reference(ws, min_col=2, min_row=32, max_row=31 + wd["n"]), title="Teams"); s.marker.symbol = "circle"; s.graphicalProperties.line.noFill = True
    sc.series.append(s); ws.add_chart(sc, "F22")
    # title odds helper
    tb = [r for r in range(sm["first"], sm["last"]) if False]
    put(ws, "A65", "Title odds: model vs market (fair)", "text", bold=True); headers(ws, 66, ["Team", "Market fair", "Model"])
    tm = o[o.market == "TITLE"].sort_values("fair", ascending=False)
    for i, (idx, x) in enumerate(tm.iterrows()):
        r = 67 + i; sr = sm["rowmap"][idx]
        put(ws, f"A{r}", f"=S_Markets!C{sr}", "link"); put(ws, f"B{r}", f"=S_Markets!I{sr}", "link", "pct"); put(ws, f"C{r}", f"=S_Markets!J{sr}", "link", "pct")
    bc = BarChart(); bc.title = "Title probability: market fair vs model"; bc.height = 9; bc.width = 18
    bc.add_data(Reference(ws, min_col=2, max_col=3, min_row=66, max_row=66 + len(tm)), titles_from_data=True); bc.set_categories(Reference(ws, min_col=1, min_row=67, max_row=66 + len(tm))); ws.add_chart(bc, "F40")
    widths(ws, {"A": 34, "B": 44}); finish(ws)


def build_sensitivity(wb):
    ws = wb.create_sheet("Sensitivity"); sheet_header(ws, "Sensitivity - title / wins probabilities vs key Assumptions (engine reruns) + scenario selector", span=10)
    put(ws, "A5", "Scenario selector lives on Assumptions!B5 (Bull / Base / Bear). It reprices the analytic layer (M_TeamRatings, M_WinDist) live; sim-based probabilities need `run.py sim`.", "text", italic=True)
    D = pd.read_csv(C("sim_sensitivity.csv")); top = D[D.case == "base"].sort_values("p_title", ascending=False).head(6).team.tolist()
    cases = list(dict.fromkeys(D.case)); headers(ws, 7, ["Case"] + [f"{t} title" for t in top] + ["Max |wins change|"], height=30)
    W = D.pivot(index="team", columns="case", values="wins"); P = D.pivot(index="team", columns="case", values="p_title")
    for i, c in enumerate(cases):
        r = 8 + i; put(ws, f"A{r}", c, "text")
        for j, t in enumerate(top): put(ws, f"{CL(2+j)}{r}", float(P.loc[t, c]), "input", "pct")
        put(ws, f"{CL(2+len(top))}{r}", float((W[c] - W["base"]).abs().max()), "input", "num2")
    last = 7 + len(cases)
    put(ws, f"A{last+2}", "Tornado: swing in title probability (max - min across cases) per team", "text", bold=True); headers(ws, last + 3, ["Team", "Low", "High", "Swing"])
    for j, t in enumerate(top):
        r = last + 4 + j; put(ws, f"A{r}", t, "text"); put(ws, f"B{r}", float(P.loc[t].min()), "input", "pct"); put(ws, f"C{r}", float(P.loc[t].max()), "input", "pct"); put(ws, f"D{r}", f"=C{r}-B{r}", "formula", "pct")
    bc = BarChart(); bc.type = "bar"; bc.title = "Title probability swing across assumption cases"; bc.height = 8; bc.width = 16
    bc.add_data(Reference(ws, min_col=4, min_row=last + 3, max_row=last + 3 + len(top)), titles_from_data=True); bc.set_categories(Reference(ws, min_col=1, min_row=last + 4, max_row=last + 3 + len(top))); ws.add_chart(bc, "J7")
    LAC = D[D.case == "LAC+TOR own-pick toggle"].set_index("team")
    widths(ws, {"A": 34}); finish(ws)


def build_backtest(wb):
    ws = wb.create_sheet("Backtest"); sheet_header(ws, "Backtest - G2 ratings backtest and G3 simulation calibration (train/test marked)", span=8)
    bt = json.load(open(C("backtest_metrics_firstroster.json"))); cal = json.load(open(C("sim_calibration.json"))); cv = json.load(open(C("sigma_team_cv_meanmatched.json")))
    put(ws, "A5", "Train/test: each test season t (2018-19..2025-26) is predicted using only seasons < t (expanding window). Headline = first-team-roster point-in-time proxy (RMSE 8.92); actual-roster (leaky, 8.44) and prior-roster (no moves, 10.30) bound it.", "text", italic=True)
    headers(ws, 7, ["Model", "Wins RMSE", "RMSE CI low (80%)", "RMSE CI high (80%)", "Wins MAE", "Spearman rho", "Paired RMSE diff vs model (baseline-model)"])
    names = {"model": "This model (first-team roster)", "B1_prev_wins_regress": "B1 prior-wins regression", "B2_prev_srs": "B2 prior-SRS regression", "B3_naive_prev_wins": "B3 naive prior wins", "B4_prev_srs_2of3_continuity": "B4 prior SRS 2/3 + continuity"}
    for i, (k, nm) in enumerate(names.items()):
        v = bt[k]; r = 8 + i; put(ws, f"A{r}", nm, "text")
        for col, x, f in (("B", v["rmse"], "num2"), ("C", v["rmse_ci"][0], "num2"), ("D", v["rmse_ci"][1], "num2"), ("E", v["mae"], "num2"), ("F", v["rho"], "num3")): put(ws, f"{col}{r}", float(x), "input", f)
        d = v.get("paired_rmse_diff_model_minus_this"); put(ws, f"G{r}", (-float(d["mean"])) if d else "-", "input", "num2")
    put(ws, "A15", "Season RMSE (model vs B1 vs B2 vs B4)", "text", bold=True); headers(ws, 16, ["Test season", "Model", "B1", "B2", "B4"])
    for i, (s, v) in enumerate(bt["by_season"].items()):
        r = 17 + i; put(ws, f"A{r}", int(s), "input", "int")
        for col, k in zip("BCDE", ("model", "B1", "B2", "B4")): put(ws, f"{col}{r}", float(v[k]), "input", "num2")
    r0 = 27; put(ws, f"A{r0}", "G3 simulation calibration (2023-24..2025-26 preseason inputs; sigma_team via leave-one-season-out CV)", "text", bold=True)
    kv = [("Team-seasons", cal["n_team_seasons"]), ("Interval coverage 50%", cal["interval_coverage"]["50%"]), ("Interval coverage 80%", cal["interval_coverage"]["80%"]), ("Interval coverage 90%", cal["interval_coverage"]["90%"]),
          ("Variance ratio observed", cal["variance_ratio_obs"]), ("Ceiling mean", cal["variance_ratio_ceiling"]["mean"]), ("Ceiling 95th pct", cal["variance_ratio_ceiling"]["p95"]), ("sigma_team CV optimum (SRS pts)", cv["sigma_all_season_optimum"])]
    for i, (k, v) in enumerate(kv): put(ws, f"A{r0+1+i}", k, "text"); put(ws, f"B{r0+1+i}", float(v), "input", "num3")
    r1 = r0 + 11; put(ws, f"A{r1}", "Reliability: P(top-10 in conference)", "text", bold=True); headers(ws, r1 + 1, ["Bin", "n", "Predicted", "Observed"])
    for i, x in enumerate(cal["calibration"]["P(top-10 in conference = makes playoffs/play-in)"]["reliability"]):
        rr = r1 + 2 + i; put(ws, f"A{rr}", x["bin"], "input"); put(ws, f"B{rr}", x["n"], "input", "int"); put(ws, f"C{rr}", x["mean_pred"], "input", "pct"); put(ws, f"D{rr}", x["freq"], "input", "pct")
    ch = ScatterChart(); ch.title = "Reliability (P(top-10))"; ch.style = 13; ch.height = 8; ch.width = 12
    n = len(cal["calibration"]["P(top-10 in conference = makes playoffs/play-in)"]["reliability"])
    s = Series(Reference(ws, min_col=4, min_row=r1 + 2, max_row=r1 + 1 + n), Reference(ws, min_col=3, min_row=r1 + 2, max_row=r1 + 1 + n), title="Observed"); s.marker.symbol = "circle"; ch.series.append(s); ws.add_chart(ch, "J7")
    widths(ws, {"A": 40}); finish(ws)


def build_tracker(wb):
    ws = wb.create_sheet("Tracker"); sheet_header(ws, "Tracker - placed bets (blue = inputs); CLV = closing implied - placed implied", span=12)
    headers(ws, 7, ["Date placed", "Market", "Selection", "Side", "Line", "Placed American", "Stake $", "Closing American", "CLV (implied pp)", "Settled (Y/N)", "Result (W/L/P)", "P&L $", "ROI on stake"])
    for r in range(8, 108):
        for col, f in (("A", "date"), ("B", None), ("C", None), ("D", None), ("E", "num1"), ("F", "amer"), ("G", "money"), ("H", "amer"), ("J", None), ("K", None)):
            put(ws, f"{col}{r}", None, "input", f)
        put(ws, f"I{r}", f'=IF(OR(F{r}="",H{r}=""),"-",(IF(H{r}>0,100/(H{r}+100),-H{r}/(-H{r}+100))-IF(F{r}>0,100/(F{r}+100),-F{r}/(-F{r}+100)))*100)', "formula", "num2")
        put(ws, f"L{r}", f'=IF(OR(G{r}="",K{r}=""),"-",IF(K{r}="W",G{r}*IF(F{r}>0,F{r}/100,100/-F{r}),IF(K{r}="L",-G{r},0)))', "formula", "money")
        put(ws, f"M{r}", f'=IF(OR(L{r}="-",G{r}=""),"-",L{r}/G{r})', "formula", "pct")
    dv = DataValidation(type="list", formula1='"W,L,P"', allow_blank=True); ws.add_data_validation(dv); dv.add("K8:K107")
    put(ws, "O7", "ROI by market type", "header"); put(ws, "P7", "Stake $", "header"); put(ws, "Q7", "P&L $", "header"); put(ws, "R7", "ROI", "header")
    for i, m in enumerate(["TITLE", "CONF", "DIV", "CUP", "PLAYIN", "PLAYOFFS", "WINS", "AWARDS/LEADERS"]):
        r = 8 + i; put(ws, f"O{r}", m, "text")
        if m == "AWARDS/LEADERS":
            put(ws, f"P{r}", '=SUM(G8:G107)-SUM(P8:P14)', "formula", "money"); put(ws, f"Q{r}", '=SUMPRODUCT(--ISNUMBER(L8:L107),L8:L107)-SUM(Q8:Q14)', "formula", "money")
        else:
            put(ws, f"P{r}", f'=SUMIF(B8:B107,"{m}*",G8:G107)', "formula", "money"); put(ws, f"Q{r}", f'=SUMPRODUCT((LEFT($B$8:$B$107,{len(m)})="{m}")*ISNUMBER(L8:L107),L8:L107)', "formula", "money")
        put(ws, f"R{r}", f'=IF(P{r}=0,"-",Q{r}/P{r})', "formula", "pct")
    widths(ws, {"A": 14, "B": 16, "C": 28, "O": 20}); finish(ws)


def build_checks(wb, o):
    ws = wb.create_sheet("Checks"); sheet_header(ws, "Checks - integrity flags (green 0 = pass; red = delta)", span=6)
    headers(ws, 7, ["Check", "Value / delta", "Status"])
    dv = REG["devig"]; od = REG["odds"]; sm = REG["sims_team"]; cm = sm["cols"]; tr = REG["tr"]; f, l = sm["first"], sm["last"]
    A = lambda k: f"M_Sims!${cm[k]}${f}:${cm[k]}${l}"
    chk = []
    # fair probabilities sum per market (max abs deviation from target S)
    mk = sorted(set(o.market) - {"WINS", "PLAYOFFS"})
    def fair_dev(m):
        return f'ABS(SUMIF(M_DeVig!$B${dv["first"]}:$B${dv["last"]},"{m}",M_DeVig!$Q${dv["first"]}:$Q${dv["last"]})-IF(LEFT("{m}",6)="PLAYIN",A_playin_slots,1))'
    chk.append(("Fair probabilities sum to target per market (max deviation, multi-outcome markets)", "=MAX(" + ",".join(fair_dev(m) for m in mk) + ")"))
    chk.append(("Pair markets (WINS/PLAYOFFS): pairs whose fair probabilities do not sum to 1", f'=SUMPRODUCT(((M_DeVig!$B${dv["first"]}:$B${dv["last"]}="WINS")+(M_DeVig!$B${dv["first"]}:$B${dv["last"]}="PLAYOFFS"))*(ABS(SUMIF(M_DeVig!$G${dv["first"]}:$G${dv["last"]},M_DeVig!$G${dv["first"]}:$G${dv["last"]},M_DeVig!$Q${dv["first"]}:$Q${dv["last"]})-1)>0.000001))'))
    chk.append(("Model probabilities sum to 1: title", f"=ABS(SUM({A('p_title')})-1)")); chk.append(("Model probabilities sum to 2: conference", f"=ABS(SUM({A('p_conference')})-2)"))
    chk.append(("Model probabilities sum to 6: division", f"=ABS(SUM({A('p_division')})-6)")); chk.append(("Model probabilities sum to 1: Cup", f"=ABS(SUM({A('p_cup')})-1)"))
    chk.append(("Model probabilities sum to 6: Cup groups", f"=ABS(SUM({A('p_cup_group')})-6)"))
    chk.append(("Model wins sum to 1,230 (live M_TeamRatings)", f"=ABS(M_TeamRatings!K{tr['total']}-A_leaguew)"))
    chk.append(("P(top-10) >= P(top-6): violations", f"=SUMPRODUCT(--({A('p_top6')}>{A('p_playoffs')}+0.000001))"))
    chk.append(("P(conference) <= P(top-10): violations", f"=SUMPRODUCT(--({A('p_conference')}>{A('p_playoffs')}+0.000001))"))
    chk.append(("P(title) <= P(conference): violations", f"=SUMPRODUCT(--({A('p_title')}>{A('p_conference')}+0.000001))"))
    put(ws, "E7", "Expected odds rows", "header"); put(ws, "E8", od["n"], "input", "int")
    chk.append(("Odds file row count reconciles (rows - expected in E8)", f"=COUNTA(Odds_Input!$B${od['first']}:$B${od['last']})-$E$8"))
    chk.append(("Imputed-odds rows (info; never bet)", f'=COUNTIF(Odds_Input!$G${od["first"]}:$G${od["last"]},"IMPUTED_20C_LADDER")'))
    chk.append(("Backfilled-odds rows (info)", f'=COUNTIF(Odds_Input!$G${od["first"]}:$G${od["last"]},"BACKFILL*")'))
    chk.append(("Live analytic wins vs Python reference (max |delta| wins)", f"=MAX(INDEX(ABS(M_TeamRatings!T{tr['first']}:T{tr['last']}),0))"))
    chk.append(("Analytic vs sim P(over) deltas > 2pp (count)", f'=COUNTIF(M_WinDist!I{REG["windist"]["first"]}:I{REG["windist"]["last"]},"CHECK*")'))
    chk.append(("Sim assumptions key matches current Assumptions (0 = match)", '=IF(A_HashCurrent=M_Sims!B9,0,1)'))
    smk_ = REG["smk"]
    chk.append(("Live Tier 1 count vs engine Tier 1 count (0 = match; differs after changing thresholds)", f'=COUNTIF(S_Markets!$AH${smk_["first"]}:$AH${smk_["last"]},TRUE)+COUNTIF(M_Players!$O${REG["m_players"]["first"]}:$O${REG["m_players"]["last"]},TRUE)-M_Sims!B12-M_Players!$B$6'))
    chk.append(("Live Tier 2 count vs engine Tier 2 count (0 = match)", f'=COUNTIF(S_Markets!$W${smk_["first"]}:$W${smk_["last"]},TRUE)-M_Sims!D12'))
    chk.append(("Tier 1 total stake within cap (excess over A_t1_cap, 0 = OK)", f'=MAX(0,Exposure!B14+Exposure!B15-A_t1_cap)'))
    for sh, rng in (("M_DeVig", f"A8:S{dv['last']}"), ("M_TeamRatings", f"A8:T{tr['last']}"), ("M_WinDist", f"A8:N{REG['windist']['last']}"), ("S_Markets", f"A8:AA{REG['smk']['last']}"), ("Bet_Card", "A8:P80"), ("Tier1_Card", "A8:N60"), ("Exposure", "A8:C58"), ("Dashboard", "A6:D80")):
        chk.append((f"Error cells in {sh} (#N/A/#VALUE!/#REF!/...)", f"=SUMPRODUCT(--ISERROR({sh}!{rng}))"))
    for i, (k, fm) in enumerate(chk):
        r = 8 + i; put(ws, f"A{r}", k, "text"); put(ws, f"B{r}", fm, "formula", "0.0000"); put(ws, f"C{r}", f'=IF(ABS(B{r})<0.0001,"OK","CHECK")', "formula")
        if "info" in k: ws[f"C{r}"].value = '="INFO"'
        if "Analytic vs sim" in k: ws[f"C{r}"].value = f'=IF(B{r}=0,"OK","FLAG")'
    last = 7 + len(chk)
    ws.conditional_formatting.add(f"B8:B{last}", CellIsRule(operator="lessThan", formula=["0.0001"], font=Font(name="Arial", size=10, color="008000", bold=True)))
    ws.conditional_formatting.add(f"B8:B{last}", CellIsRule(operator="greaterThanOrEqual", formula=["0.0001"], font=Font(name="Arial", size=10, color=RED, bold=True)))
    put(ws, f"A{last+2}", "Overall", "text", bold=True); put(ws, f"B{last+2}", f'=COUNTIF(C8:C{last},"CHECK")', "formula", "int"); total_border(ws, last + 2, 1, 3)
    REG["checks"] = dict(first=8, last=last, overall=last + 2)
    widths(ws, {"A": 82, "B": 16, "C": 12}); finish(ws)


def build_sources(wb):
    ws = wb.create_sheet("Data_Sources"); sheet_header(ws, "Data_Sources - every source: URL, pull timestamp, rows, status, license note", span=7)
    headers(ws, 7, ["Source", "URL / endpoint", "Pull timestamp", "Rows", "Status", "Fallback used", "License / note"])
    M = pd.read_csv(os.path.join(ROOT, "data/manifest.csv"))
    M["grp"] = M.file.str.extract(r"data/(?:raw|clean)/([^/]+)")[0].fillna("params")
    agg = M.groupby("grp").agg(url=("source_url", "first"), ts=("pull_timestamp", "max"), rows=("row_count", "sum"), n=("file", "count")).reset_index()
    note = {"espn": "ESPN public site API - rosters/schedule/injuries (nba_api blocked)", "schedule": "bbref schedule pages (cdn.nba.com 403)", "player_rs": "bbref (nba_api blocked)", "player_po": "bbref", "team": "bbref", "games": "bbref", "draft": "bbref",
            "awards": "bbref", "darko": "darko.app leaderboard read via browser pane (CSV hash-verified)", "barttorvik": "barttorvik.com CSV (headerless)", "odds": "user-supplied", "lottery": "NBA.com/ESPN/Hoops Rumors (text summaries)", "picks": "Hoops Rumors / RealGM / Tankathon"}
    for i, r in agg.iterrows():
        row = 8 + i
        put(ws, f"A{row}", r.grp, "input"); put(ws, f"B{row}", str(r.url)[:120], "input"); put(ws, f"C{row}", str(r.ts), "input"); put(ws, f"D{row}", int(r.rows) if pd.notna(r.rows) else 0, "input", "int")
        put(ws, f"E{row}", "OK", "input"); put(ws, f"F{row}", "bbref/ESPN used because stats.nba.com / cdn.nba.com are blocked from this network" if r.grp in ("espn", "schedule", "player_rs", "player_po", "team", "games") else "-", "input"); put(ws, f"G{row}", note.get(r.grp, "derived / model output"), "input")
    widths(ws, {"A": 20, "B": 60, "C": 28, "F": 50, "G": 60}); finish(ws)


def build_cover(wb, key):
    ws = wb.create_sheet("Cover", 0); sheet_header(ws, "Cover", span=6)
    meta = json.load(open(C("sim_run_meta.json")))
    put(ws, "A5", "NBA 2026-27 Futures Model", "text", bold=True); ws["A5"].font = font(NAVY, bold=True, size=16)
    rows = [("Version", "v1 (G5) - 2026-10-15"), ("Data timestamp", "Odds 2026-10-15; rosters/injuries ESPN 2026-10-15; DARKO through 2026-06-03"), ("Sim N", "=M_Sims!B6"), ("Run hash", "=M_Sims!B8"), ("Run timestamp (UTC)", "=M_Sims!B5"), ("Sim status", "=Sim_Stale"),
            ("Model independence", "Market odds are never a model input; they enter only de-vig, edges, staking shrink and backtest benchmarks."), ("First locks", "Tue 10/20 2:10 PM CST (BOS/DET wins & playoffs; all title, conference, division, play-in, award markets)")]
    for i, (k, v) in enumerate(rows):
        put(ws, f"A{7+i}", k, "text", bold=True); put(ws, f"B{7+i}", v, "link" if str(v).startswith("=") else "text")
    put(ws, "A17", "Navigation", "text", bold=True)
    tabs = ["Assumptions", "Overrides", "Odds_Input", "M_DeVig", "M_TeamRatings", "M_WinDist", "M_Sims", "M_Lottery2027", "M_Players", "S_Markets", "Bet_Card", "Tier1_Card", "Exposure", "Dashboard", "Sensitivity", "Backtest", "Tracker", "Checks", "Data_Sources", "Raw_Teams", "Raw_Players", "Raw_Schedule"]
    for i, t in enumerate(tabs):
        put(ws, f"A{18+i}", f'=HYPERLINK("#\'{t}\'!A1","{t}")', "formula"); ws[f"A{18+i}"].font = Font(name="Arial", size=10, color="0563C1", underline="single")
    put(ws, "A41", "Open items: PLAYOFFS = 16-team playoff field (user-confirmed from implied sums) - CONFIRM WITH BOOK. Tier 1 (Tier1_Card) is the staked card; Tier 2 (Bet_Card) has no stake by default. Player markets: see M_Players.", "flag", wrap=False)
    widths(ws, {"A": 30, "B": 100}); finish(ws, label_cols=("A", "B"))
